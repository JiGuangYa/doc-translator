"""xlsx preview regression: streaming iter_rows render structure, truncation,
and None-dimension fallback."""
import openpyxl
import pytest

from app.formats import xlsx_preview


def _make_book(path, rows, cols):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(row=r, column=c, value=f"r{r}c{c}")
    wb.save(path)


@pytest.fixture()
def small_xlsx(tmp_path):
    p = tmp_path / "small.xlsx"
    _make_book(p, rows=5, cols=3)
    return p


def test_render_contains_cells(small_xlsx):
    html = xlsx_preview.render_xlsx_dual_html(small_xlsx, small_xlsx)
    assert "r1c1" in html and "r5c3" in html
    # Two panes: cell text appears in the title attribute and in the cell body,
    # for source + translation panes that is 4 occurrences total
    assert html.count("r1c1") == 4
    assert ">Original" in html and ">Translated" in html
    assert "xl-tab" in html


def test_truncation_note_appears(tmp_path):
    p = tmp_path / "big.xlsx"
    _make_book(p, rows=xlsx_preview._MAX_ROWS + 10, cols=2)
    html = xlsx_preview.render_xlsx_dual_html(p, p)
    assert "Showing only the first" in html
    # After truncation, out-of-range rows must not appear
    assert f"r{xlsx_preview._MAX_ROWS + 1}c1" not in html


def test_none_dimension_fallback(tmp_path):
    """In read_only mode a missing dimension leaves max_row/max_column as
    None, and must not raise TypeError."""
    p = tmp_path / "nodim.xlsx"
    _make_book(p, rows=3, cols=2)
    wb = openpyxl.load_workbook(str(p), read_only=True)
    try:

        class NoDimSheet:
            """Simulate a read_only worksheet missing the dimension record."""

            def __init__(self, real):
                self._real = real
                self.max_row = None
                self.max_column = None

            def iter_rows(self, **kw):
                return self._real.iter_rows(**kw)

        table = xlsx_preview._render_sheet_table(NoDimSheet(wb["Sheet1"]), "Original")
        assert "r3c2" in table
    finally:
        wb.close()
