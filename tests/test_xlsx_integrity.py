"""Complex workbook translation must preserve non-text Office parts."""

import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as WorkbookImage
from PIL import Image

from app.formats import xlsx_fmt
from app.formats.common import FormatAdapterError
from app.services import renderer


def _complex_workbook(tmp_path):
    source = tmp_path / "source.xlsx"
    picture = tmp_path / "picture.png"
    Image.new("RGB", (12, 12), "red").save(picture)
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Data"
    sheet["A1"] = "Revenue by region"
    sheet["A2"] = "North region"
    sheet["B1"] = "Revenue"
    sheet["B2"] = 42
    sheet["B3"] = "=B2*2"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "D2")
    sheet.add_image(WorkbookImage(picture), "F2")
    wb.save(source)
    return source


def test_without_genoffice_complex_workbook_is_blocked(tmp_path, monkeypatch):
    source = _complex_workbook(tmp_path)
    monkeypatch.setattr(renderer, "genoffice_path", lambda: None)
    with pytest.raises(FormatAdapterError, match="Install GenOffice"):
        xlsx_fmt.write_back(source, tmp_path / "translated.xlsx", {"s000001": "北部地区"}, {})


def test_genoffice_keeps_chart_image_and_formula(tmp_path):
    if not renderer.genoffice_path():
        pytest.skip("GenOffice CLI is not installed")
    source = _complex_workbook(tmp_path)
    translated = tmp_path / "translated.xlsx"
    segments = xlsx_fmt.extract(source, {}).segments
    index = next(index for index, segment in enumerate(segments)
                 if segment.context == "Data!A2")
    report = xlsx_fmt.write_back(source, translated,
                                 {f"s{index:06d}": "北部地区"}, {})
    assert report.written == 1
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(translated) as after:
        for name in ("xl/charts/chart1.xml", "xl/media/image1.png"):
            assert before.read(name) == after.read(name)
    wb = load_workbook(translated)
    assert wb["Data"]["A2"].value == "北部地区"
    assert wb["Data"]["B3"].value == "=B2*2"
