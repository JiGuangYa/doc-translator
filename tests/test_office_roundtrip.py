"""Round-trip tests for three Office formats (docx/pptx/xlsx) extract/write_back.

Builds small fixtures dynamically for self-containment; also includes
integration tests against real office/ samples (skipped if missing).
"""
import pytest
from pathlib import Path

from app.formats import docx_fmt, pptx_fmt, xlsx_fmt
from app.formats.xlsx_preview import render_xlsx_dual_html

OFFICE_DIR = Path(r"C:\fcc\office")


# ---------- fixtures ----------

@pytest.fixture
def tmp_paths(tmp_path):
    return {"src": tmp_path / "src", "dst": tmp_path / "dst"}


def _make_docx(path: Path) -> None:
    """Body segments + bold mixed layout + hyperlink + table + header."""
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    p1 = doc.add_paragraph()
    r = p1.add_run("Hello ")
    r.bold = True
    p1.add_run("brave world")

    # Hyperlink segment (python-docx has no high-level API, build w:hyperlink manually)
    p2 = doc.add_paragraph("See ")
    link_id = p2.part.relate_to("https://example.com/docs", RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), link_id)
    run_el = OxmlElement("w:r")
    t_el = OxmlElement("w:t")
    t_el.text = "the official manual page"
    run_el.append(t_el)
    hl.append(run_el)
    p2._p.append(hl)

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell alpha"
    table.cell(0, 1).text = "Cell beta text"
    table.cell(1, 0).text = "42"  # not translatable, but should still appear in extract and be filtered by the pipeline

    header = doc.sections[0].header
    header.paragraphs[0].add_run("Header note for all pages")

    doc.save(str(path))


def _make_pptx(path: Path) -> None:
    """Plain text box + text box inside a group shape + table."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1))
    tf = box.text_frame
    tf.text = "Title on plain slide"
    para = tf.paragraphs[0]
    run = para.add_run()
    run.text = " with mixed bold part"
    run.font.bold = True

    # Soft line break (a:br) segment: verify \x0b is not written out as broken text
    br_box = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(4), Inches(1))
    br_para = br_box.text_frame.paragraphs[0]
    br_run = br_para.add_run()
    br_run.text = "Soft break first"
    br_para.add_line_break()
    br_run2 = br_para.add_run()
    br_run2.text = "second"

    group = slide.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    inner.text_frame.text = "Text inside grouped shape"

    tbl = slide.shapes.add_table(2, 2, Inches(5), Inches(0.5), Inches(4), Inches(1.5))
    tbl.table.cell(0, 0).text = "Table head label"
    tbl.table.cell(1, 1).text = "Table body content"

    prs.save(str(path))


def _make_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = "Welcome to the report"
    ws1["B3"] = "Another translatable string"
    ws1["C1"] = 12345            # number is left untouched
    ws1["D1"] = "=SUM(C1, 1)"    # formula is left untouched
    ws2 = wb.create_sheet("Data")
    ws2["A1"] = "Second sheet value"
    wb.save(str(path))


def _identity_roundtrip(handler, src: Path, dst: Path):
    result = handler.extract(src, {})
    translations = {seg.seg_id if seg.seg_id else f"s{i:06d}": seg.text
                    for i, seg in enumerate(result.segments)}
    handler.write_back(src, dst, translations, {})
    after = handler.extract(dst, {})
    assert [s.text for s in after.segments] == [s.text for s in result.segments]
    return result, after


# ---------- docx ----------

def test_docx_roundtrip(tmp_path):
    src = tmp_path / "in.docx"
    dst = tmp_path / "out.docx"
    _make_docx(src)
    result, _after = _identity_roundtrip(docx_fmt, src, dst)

    texts = [s.text for s in result.segments]
    assert "Hello brave world" in texts          # mixed bold run concatenation
    assert "See the official manual page" in texts  # run inside hyperlink
    assert "Cell alpha" in texts and "Cell beta text" in texts
    assert "Header note for all pages" in texts  # header
    contexts = {s.text: s.context for s in result.segments}
    assert contexts["Cell alpha"] == "Table"
    assert contexts["Header note for all pages"] == "Header"
    # The output file must be re-openable by python-docx
    from docx import Document
    Document(str(dst))


def test_docx_translate_writeback(tmp_path):
    src = tmp_path / "in.docx"
    dst = tmp_path / "out.docx"
    _make_docx(src)
    result = docx_fmt.extract(src, {})
    translations = {f"s{i:06d}": f"【T】{seg.text}" for i, seg in enumerate(result.segments)}
    report = docx_fmt.write_back(src, dst, translations, {})
    assert report.written == len(result.segments)

    after = docx_fmt.extract(dst, {})
    assert [s.text for s in after.segments] == [f"【T】{s.text}" for s in result.segments]


def test_docx_real_sample(tmp_path):
    """Integration test with a real sample: after the round-trip the set of
    segment text must be identical."""
    sample = OFFICE_DIR / "CarWash_Manual_EN.docx"
    if not sample.exists():
        pytest.skip(f"real sample not found: {sample}")
    dst = tmp_path / "out.docx"
    _identity_roundtrip(docx_fmt, sample, dst)


# ---------- pptx ----------

def test_pptx_roundtrip(tmp_path):
    src = tmp_path / "in.pptx"
    dst = tmp_path / "out.pptx"
    _make_pptx(src)
    result, _after = _identity_roundtrip(pptx_fmt, src, dst)

    texts = [s.text for s in result.segments]
    assert "Title on plain slide with mixed bold part" in texts
    assert "Soft break first\x0bsecond" in texts  # a:br is extracted as \x0b and round-trips losslessly
    assert "Text inside grouped shape" in texts   # grouped shape
    assert "Table head label" in texts and "Table body content" in texts
    by_text = {s.text: s.context for s in result.segments}
    assert by_text["Text inside grouped shape"].startswith("Slide 1 / group shape")
    assert "/ table R2C2" in by_text["Table body content"]


def test_pptx_writeback(tmp_path):
    src = tmp_path / "in.pptx"
    dst = tmp_path / "out.pptx"
    _make_pptx(src)
    result = pptx_fmt.extract(src, {})
    translations = {f"s{s_i:06d}": f"【T】{seg.text}" for s_i, seg in enumerate(result.segments)}
    options = {"target_lang": "zh-CN"}
    report = pptx_fmt.write_back(src, dst, translations, options)
    assert report.written == len(result.segments)

    after = pptx_fmt.extract(dst, {})
    assert [s.text for s in after.segments] == [f"【T】{s.text}" for s in result.segments]

    # lang attribute is written in the target language (check the run that carries the translation)
    from pptx import Presentation
    prs = Presentation(str(dst))
    found = False
    for shape in prs.slides[0].shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.text.startswith("【T】"):
                    rpr = run._r.find(
                        "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                    assert rpr is not None and rpr.get("lang") == "zh-CN"
                    found = True
    assert found


def test_pptx_real_sample(tmp_path):
    sample = OFFICE_DIR / "CarWash_Training_EN.pptx"
    if not sample.exists():
        pytest.skip(f"real sample not found: {sample}")
    dst = tmp_path / "out.pptx"
    _identity_roundtrip(pptx_fmt, sample, dst)


# ---------- xlsx ----------

def test_xlsx_roundtrip(tmp_path):
    src = tmp_path / "in.xlsx"
    dst = tmp_path / "out.xlsx"
    _make_xlsx(src)
    result = xlsx_fmt.extract(src, {})
    texts = {s.text: s.context for s in result.segments}
    assert "Welcome to the report" in texts and texts["Welcome to the report"] == "Sheet1!A1"
    assert "Second sheet value" in texts and texts["Second sheet value"] == "Data!A1"

    translations = {f"s{i:06d}": seg.text for i, seg in enumerate(result.segments)}
    xlsx_fmt.write_back(src, dst, translations, {})

    import openpyxl
    wb = openpyxl.load_workbook(str(dst))
    ws1, ws2 = wb["Sheet1"], wb["Data"]
    assert ws1["A1"].value == "Welcome to the report"
    assert ws1["C1"].value == 12345                       # number is unchanged
    assert str(ws1["D1"].value).replace("$", "").upper().startswith("=")  # formula preserved
    assert ws2["A1"].value == "Second sheet value"


def test_xlsx_translate_only_strings(tmp_path):
    src = tmp_path / "in.xlsx"
    dst = tmp_path / "out.xlsx"
    _make_xlsx(src)
    result = xlsx_fmt.extract(src, {})
    translations = {
        f"s{i:06d}": f"[X]{seg.text}"
        for i, seg in enumerate(result.segments)
        if seg.context == "Sheet1!A1"
    }
    xlsx_fmt.write_back(src, dst, translations, {})

    import openpyxl
    wb = openpyxl.load_workbook(str(dst))
    assert wb["Sheet1"]["A1"].value == "[X]Welcome to the report"
    assert wb["Sheet1"]["B3"].value == "Another translatable string"  # untranslated, original kept
    assert wb["Sheet1"]["C1"].value == 12345


def test_xlsx_preview_html(tmp_path):
    src = tmp_path / "orig.xlsx"
    dst = tmp_path / "trans.xlsx"
    _make_xlsx(src)
    result = xlsx_fmt.extract(src, {})
    translations = {f"s{i:06d}": f"[X]{seg.text}" for i, seg in enumerate(result.segments)}
    xlsx_fmt.write_back(src, dst, translations, {})

    html_out = render_xlsx_dual_html(src, dst)
    assert "xl-dual" in html_out and "#14171c" in html_out
    assert "[X]Welcome to the report" in html_out
    assert html_out.count("xl-pane") >= 2  # source + translation two panes


def test_xlsx_real_sample(tmp_path):
    sample = OFFICE_DIR / "i2Active_RAC_Investigation_and_Prerequisites.xlsx"
    if not sample.exists():
        pytest.skip(f"real sample not found: {sample}")
    dst = tmp_path / "out.xlsx"
    _identity_roundtrip(xlsx_fmt, sample, dst)
