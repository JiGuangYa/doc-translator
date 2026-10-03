"""Clickable links stay intact and segment reading order is preserved."""

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app import config
from app.services import pipeline


def test_legacy_hyperlink_segment_is_visible_but_not_billed_for_translation(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TASKS_DIR", tmp_path / "tasks")
    source = tmp_path / "links.docx"
    document = Document()
    document.add_paragraph("First paragraph is translatable")
    paragraph = document.add_paragraph("Read ")
    relation = paragraph.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    content = OxmlElement("w:t")
    content.text = "the documentation"
    run.append(content)
    link.append(run)
    paragraph._p.append(link)
    document.add_paragraph("Last paragraph is translatable")
    document.save(source)

    job = pipeline.parse_upload("a" * 32, source.name, source, {"docx_version": 1})
    assert [segment["text"] for segment in job["segments"]] == [
        "First paragraph is translatable", "Read the documentation",
        "Last paragraph is translatable"]
    assert [segment["translatable"] for segment in job["segments"]] == [True, False, True]
