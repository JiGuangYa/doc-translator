"""Text translation cannot silently discard embedded Word or slide media."""

import zipfile

from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.formats import docx_fmt, pptx_fmt
from app.services.pipeline import _validate_output


def _picture(tmp_path):
    path = tmp_path / "picture.png"
    Image.new("RGB", (24, 24), "blue").save(path)
    return path


def test_docx_image_survives_text_writeback(tmp_path):
    picture = _picture(tmp_path)
    source = tmp_path / "source.docx"
    target = tmp_path / "target.docx"
    doc = Document()
    doc.add_paragraph("Translate this document")
    doc.add_picture(str(picture))
    doc.save(source)
    docx_fmt.write_back(source, target, {"s000000": "翻译这份文档"}, {})
    _validate_output(target, ".docx", source)
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(target) as after:
        assert before.read("word/media/image1.png") == after.read("word/media/image1.png")


def test_pptx_image_survives_text_writeback(tmp_path):
    picture = _picture(tmp_path)
    source = tmp_path / "source.pptx"
    target = tmp_path / "target.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "Translate slide text"
    slide.shapes.add_picture(str(picture), Inches(1), Inches(3))
    presentation.save(source)
    pptx_fmt.write_back(source, target, {"s000000": "翻译幻灯片文字"}, {})
    _validate_output(target, ".pptx", source)
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(target) as after:
        assert before.read("ppt/media/image1.png") == after.read("ppt/media/image1.png")
