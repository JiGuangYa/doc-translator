"""Long translations in slides surface a visual review warning."""

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.formats import pptx_fmt
from app.services.renderer import genoffice_path


def test_overflow_warning_from_genoffice(tmp_path):
    if not genoffice_path():
        pytest.skip("GenOffice CLI is not installed")
    original = tmp_path / "original.pptx"
    translated = tmp_path / "translated.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.4))
    box.text = "Short title"
    presentation.save(original)
    report = pptx_fmt.write_back(original, translated, {
        "s000000": "This is a very long translation that cannot fit in the small text box " * 3
    }, {})
    assert report.written == 1
    assert any("overflow" in warning for warning in report.warnings)
