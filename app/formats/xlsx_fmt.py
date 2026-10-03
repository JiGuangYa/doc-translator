"""XLSX extraction and safe text-only write-back.

GenOffice patches cells while preserving untouched workbook parts. The
openpyxl fallback is limited to workbooks without complex embedded objects.
"""
import json
import shutil
import zipfile
from pathlib import Path

from .common import ExtractResult, FormatAdapterError, Segment, WriteReport

_PROTECTED_PARTS = (
    "xl/charts/", "xl/media/", "xl/drawings/", "xl/pivot", "xl/slicer",
    "xl/externalLinks/", "xl/embeddings/", "xl/activeX/", "xl/printerSettings/",
    "customXml/", "_xmlsignatures/",
)


def _is_skippable(cell) -> bool:
    """Skip formulas, non-strings (numbers, dates, booleans)."""
    if cell.data_type == "f":
        return True
    if not isinstance(cell.value, str):
        return True
    return cell.value.startswith("=")


def _iter_translatable_cells(wb) -> tuple[list, int]:
    """Return [(cell, context)] in iteration order. Only visible sheets; untranslatable
    strings are counted as skipped."""
    from ..utils import is_translatable

    items: list = []
    skipped = 0
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        for row in ws.iter_rows():
            for cell in row:
                if _is_skippable(cell):
                    continue
                ok, _reason = is_translatable(cell.value)
                if not ok:
                    skipped += 1
                    continue
                items.append((cell, f"{ws.title}!{cell.coordinate}"))
    return items, skipped


def extract(path: Path, options: dict) -> ExtractResult:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, keep_links=False)
    cells, skipped = _iter_translatable_cells(wb)
    segments = [Segment(seg_id="", text=cell.value, context=ctx) for cell, ctx in cells]
    wb.close()
    return ExtractResult(segments=segments, skipped_count=skipped, warnings=[])


def _protected_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {name for name in archive.namelist()
                if name.startswith(_PROTECTED_PARTS) or name.endswith("vbaProject.bin")}


def _verify_protected_parts(source: Path, output: Path) -> None:
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output) as translated:
        for name in _protected_names(source):
            if name not in translated.namelist() or original.read(name) != translated.read(name):
                raise FormatAdapterError(f"XLSX embedded object changed or disappeared: {name}")


def _verify_formulas(source: Path, output: Path) -> None:
    import openpyxl
    before = openpyxl.load_workbook(source, read_only=True, data_only=False)
    after = openpyxl.load_workbook(output, read_only=True, data_only=False)
    try:
        if before.sheetnames != after.sheetnames:
            raise FormatAdapterError("XLSX worksheet names changed during write-back")
        for name in before.sheetnames:
            original = before[name]
            written = after[name]
            for row in original.iter_rows():
                for cell in row:
                    if cell.data_type == "f" and written[cell.coordinate].value != cell.value:
                        raise FormatAdapterError(
                            f"XLSX formula changed during write-back: {name}!{cell.coordinate}")
    finally:
        before.close()
        after.close()


def write_back(src_path: Path, dst_path: Path, translations: dict[str, str],
               options: dict) -> WriteReport:
    import openpyxl
    from ..services.renderer import genoffice_path
    from ..services import process_runner

    report = WriteReport()
    if not any(translations.values()):
        shutil.copy2(src_path, dst_path)
        return report

    read_book = openpyxl.load_workbook(str(src_path), read_only=True, keep_links=False)
    try:
        cells, _skipped = _iter_translatable_cells(read_book)
        changes = []
        for index, (cell, _ctx) in enumerate(cells):
            translation = translations.get(f"s{index:06d}")
            if translation:
                changes.append({"sheet": cell.parent.title, "cell": cell.coordinate,
                                "value": translation})
    finally:
        read_book.close()
    if not changes:
        shutil.copy2(src_path, dst_path)
        return report

    cli = genoffice_path()
    if cli:
        result = process_runner.run(
            [cli, "sheet", "apply", str(src_path), "--cells", "-",
             "--out", str(dst_path), "--json"],
            input_text=json.dumps(changes, ensure_ascii=False), timeout=300,
            label="GenOffice XLSX write-back",
        )
        try:
            answer = json.loads(result.stdout.strip().splitlines()[-1])
        except (ValueError, IndexError):
            answer = {}
        if result.returncode or answer.get("status") != "ok" or not dst_path.exists():
            detail = answer.get("message") or result.stderr[-300:] or result.stdout[-300:]
            raise FormatAdapterError(f"GenOffice XLSX write-back failed: {detail}")
        report.written = len(changes)
    else:
        protected = _protected_names(src_path)
        if protected:
            raise FormatAdapterError(
                "This workbook contains charts, images, or other embedded objects. "
                "Install GenOffice to translate it without losing those objects.")
        shutil.copy2(src_path, dst_path)
        wb = openpyxl.load_workbook(str(dst_path), keep_links=False)
        for change in changes:
            wb[change["sheet"]][change["cell"]] = change["value"]
        wb.save(str(dst_path))
        report.written = len(changes)
        report.warnings.append("GenOffice unavailable; simple workbook written with openpyxl")

    _verify_protected_parts(src_path, dst_path)
    _verify_formulas(src_path, dst_path)
    return report
