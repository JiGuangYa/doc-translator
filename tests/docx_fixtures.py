"""Portable Word structure fixture shared by adapter and rendered acceptance tests."""

import copy
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw

V = "urn:schemas-microsoft-com:vml"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def textbox(text):
    pict = OxmlElement("w:pict")
    rect = etree.SubElement(pict, "{" + V + "}rect", nsmap={"v": V})
    rect.set("id", "TextBoxFixture")
    rect.set("style", "width:360pt;height:55pt")
    rect.set("fillcolor", "#eef4ff")
    box = etree.SubElement(rect, "{" + V + "}textbox")
    content = etree.SubElement(box, qn("w:txbxContent"))
    paragraph = etree.SubElement(content, qn("w:p"))
    run = etree.SubElement(paragraph, qn("w:r"))
    etree.SubElement(run, qn("w:t")).text = text
    return pict


def complex_docx(path: Path, compatibility_copy=False):
    doc = Document()
    zoom = doc.settings._element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Document fidelity check", 0)
    p = doc.add_paragraph()
    run = p.add_run("Important ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("244CA1")
    p.add_run("instructions: ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to("https://example.com/manual", RT.HYPERLINK, is_external=True))
    linked_run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    linked_run.append(props)
    text = OxmlElement("w:t")
    text.text = "read the manual"
    linked_run.append(text)
    link.append(linked_run)
    p._p.append(link)
    p.add_run(" before starting.").italic = True
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Control module"
    table.cell(1, 1).text = "Ready for review"
    doc.add_paragraph("A local picture is preserved below.")
    picture = path.with_suffix(".png")
    image = Image.new("RGB", (640, 140), "#e5eefb")
    ImageDraw.Draw(image).rounded_rectangle((15, 15, 625, 125), 14, fill="#244CA1")
    image.save(picture)
    doc.add_picture(str(picture), width=Inches(4.7))
    run = doc.add_paragraph().add_run()
    pict = textbox("Text inside a nested text box")
    if compatibility_copy:
        alternate = etree.SubElement(run._r, "{" + MC + "}AlternateContent")
        choice = etree.SubElement(alternate, "{" + MC + "}Choice", Requires="w14")
        choice.append(pict)
        fallback = etree.SubElement(alternate, "{" + MC + "}Fallback")
        fallback.append(copy.deepcopy(pict))
    else:
        run._r.append(pict)
    doc.add_paragraph("A final paragraph follows the text box.")
    section.header.paragraphs[0].text = "Engineering review / Local test"
    footer = section.footer.paragraphs[0]
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    value = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    value.append(text)
    field.append(value)
    footer._p.append(field)
    doc.save(path)
    return path
