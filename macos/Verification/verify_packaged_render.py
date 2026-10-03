"""Check real GenOffice rendering with the authenticated frozen companion."""
import argparse
import json
import tempfile
import time
from pathlib import Path
from docx import Document
from packaged_api import PackagedAPI


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('app')
    parser.add_argument('--report')
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='unified-render-') as temporary:
        root = Path(temporary)
        task_id = 'a' * 32
        folder = root / 'tasks' / task_id
        folder.mkdir(parents=True)
        for variant in ('original', 'translated'):
            document = Document()
            document.add_heading('Packaged rendering test', level=1)
            document.add_paragraph('A local pagination check. ' * (1 if variant == 'original' else 600))
            document.save(folder / f'{variant}.docx')
        (folder / 'job.json').write_text(json.dumps({'task_id': task_id, 'filename': 'preview.docx',
            'ext': '.docx', 'status': 'done', 'segments': [], 'translations': {}, 'revision': 0}))
        with PackagedAPI(args.app, root) as api:
            status = api.json(f'/api/tasks/{task_id}/preview/render', {})
            deadline = time.monotonic() + 90
            while status['status'] == 'rendering' and time.monotonic() < deadline:
                time.sleep(.2)
                status = api.json(f'/api/tasks/{task_id}/preview/render')
            assert status['status'] == 'ready', status
            assert status['original_pages'] == 1 and status['translated_pages'] > 1, status
            assert api.bytes(f'/api/tasks/{task_id}/preview/render/page/{status["translated_pages"]}?variant=translated').startswith(b'\x89PNG')
            if args.report:
                Path(args.report).write_text(json.dumps(status, indent=2))
    print('PASS: real GenOffice render and unequal page counts through frozen authenticated API')


if __name__ == '__main__':
    main()
