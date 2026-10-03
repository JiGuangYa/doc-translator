"""Regression coverage for the native app's private session and revisions."""
import io
import json
import uuid

import fitz
import pytest

from app import config, desktop, store
from app.services import pipeline, renderer


def test_desktop_requires_private_token_even_with_web_session(client, monkeypatch):
    monkeypatch.setattr(desktop, '_token', 'd' * 64)
    assert client.get('/api/tasks').status_code == 401
    assert client.get('/healthz').status_code == 401
    assert client.get('/api/tasks', headers={'X-DocTranslator-Token': 'wrong'}).status_code == 401
    assert client.get('/api/tasks', headers={'X-DocTranslator-Token': 'd' * 64}).status_code == 200


def test_pdf_revision_changes_preview_generation(client):
    source = fitz.open()
    source.new_page().insert_text((50, 70), 'Hello original paragraph.')
    response = client.post('/api/upload', files={'file': ('sample.pdf', source.tobytes(), 'application/pdf')})
    source.close()
    tid = response.json()['task_id']
    job = store.load_job(tid)
    seg = job['segments'][0]['seg_id']
    pipeline._apply_translations(tid, '.pdf', {seg: 'First translation.'})
    before = client.get(f'/api/tasks/{tid}/preview/pdf/1?variant=translated')
    version = store.load_job(tid)['content_version']
    assert before.status_code == 200
    assert client.patch(f'/api/tasks/{tid}/segments/{seg}', json={'text': 'Revised translation.'}).status_code == 200
    after = client.get(f'/api/tasks/{tid}/preview/pdf/1?variant=translated')
    assert before.content != after.content
    info = client.get(f'/api/tasks/{tid}/info').json()
    assert info['content_version'] == version + 1
    assert info['original_pages'] == info['translated_pages'] == 1


def test_failed_revision_keeps_saved_job_and_file(client, monkeypatch):
    from docx import Document
    doc = Document()
    doc.add_paragraph('Hello original paragraph.')
    data = io.BytesIO()
    doc.save(data)
    tid = client.post('/api/upload', files={'file': ('sample.docx', data.getvalue())}).json()['task_id']
    seg = store.load_job(tid)['segments'][0]['seg_id']
    pipeline._apply_translations(tid, '.docx', {seg: 'Saved translation.'})
    before_job = store.load_job(tid)
    before_file = store.document_path(store.load_job(tid), 'translated').read_bytes()
    handler = pipeline.common.get_format_handler('.docx')
    def fail(*args, **kwargs):
        raise OSError('disk write failed')
    monkeypatch.setattr(handler, 'write_back', fail)
    with pytest.raises(OSError):
        pipeline.revise_segment(tid, seg, 'Unsaved draft.')
    assert store.load_job(tid) == before_job
    assert store.document_path(store.load_job(tid), 'translated').read_bytes() == before_file


def test_resume_cannot_mix_target_languages(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'TASKS_DIR', tmp_path)
    tid = uuid.uuid4().hex
    store.save_job(tid, {'task_id': tid, 'status': 'cancelled', 'provider_id': 'p', 'source_lang': 'auto',
                        'target_lang': 'zh-CN', 'translations': {'s000000': '已完成'}})
    with pytest.raises(ValueError, match='change languages'):
        pipeline.start_translation(tid, 'p', 'auto', 'en')


def test_unavailable_renderer_does_not_start_thread(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'TASKS_DIR', tmp_path)
    monkeypatch.setattr(renderer, 'available', lambda: False)
    assert renderer.start_render(uuid.uuid4().hex, '.docx')['status'] == 'unavailable'


def test_partial_settings_update_preserves_provider(client, monkeypatch, tmp_path):
    monkeypatch.setattr(config, 'SETTINGS_FILE', tmp_path / 'settings.json')
    store.save_settings({'translation_provider_id': 'p_saved', 'assistant_provider_id': 'p_assistant'})
    assert client.put('/api/settings', json={'target_lang': 'en'}).status_code == 200
    settings = store.load_settings()
    assert settings['translation_provider_id'] == 'p_saved'
    assert settings['assistant_provider_id'] == 'p_assistant'


def test_office_preview_records_both_page_counts(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'TASKS_DIR', tmp_path)
    monkeypatch.setattr(config, 'DATA_DIR', tmp_path)
    monkeypatch.setattr(renderer, 'genoffice_path', lambda: None)
    monkeypatch.setattr(renderer, 'soffice_path', lambda: 'fake-soffice')
    tid = uuid.uuid4().hex
    store.save_job(tid, {'task_id': tid, 'ext': '.docx', 'content_version': 3})
    for variant in ('original', 'translated'):
        (store.task_dir(tid) / f'{variant}.docx').write_bytes(b'test')
    def convert(soffice, src, outdir, tag, cancel_event=None):
        outdir.mkdir(parents=True, exist_ok=True)
        pdf = fitz.open()
        for _ in range(1 if src.stem == 'original' else 2):
            pdf.new_page()
        path = outdir / f'{src.stem}.pdf'
        pdf.save(path)
        pdf.close()
        return path
    monkeypatch.setattr(renderer, '_convert_to_pdf', convert)
    renderer._render_task(tid, '.docx')
    status = renderer.render_status(tid)
    assert status['original_pages'] == 1
    assert status['translated_pages'] == status['pages'] == 2
    assert renderer.page_png_path(tid, 2, 'translated') is not None
    assert renderer.page_png_path(tid, 2, 'original') is None


def test_resume_freezes_provider_and_settings_without_copying_secrets(client, monkeypatch, tmp_path):
    from docx import Document
    from app.services import task_manager
    from app.services.llm import translator
    monkeypatch.setattr(config, 'SETTINGS_FILE', tmp_path / 'settings.json')
    provider = {'id': 'p', 'name': 'Original', 'base_url': 'http://127.0.0.1:9999/v1',
                'model': 'first-model', 'api_key': 'must-not-be-copied'}
    monkeypatch.setattr(store, 'get_provider', lambda _: dict(provider))
    monkeypatch.setattr(task_manager, 'submit', lambda fn, tid: task_manager._wrap(fn, tid))
    monkeypatch.setattr(renderer, 'auto_render_after_done', lambda *args: None)
    doc = Document()
    doc.add_paragraph('First paragraph.')
    doc.add_paragraph('Second paragraph.')
    file = io.BytesIO()
    doc.save(file)
    tid = client.post('/api/upload', files={'file': ('sample.docx', file.getvalue())}).json()['task_id']
    store.save_settings({'batch_max_segments': 1, 'translate_notes': False})
    calls = []
    def translate(segments, pid, source, target, settings, **kwargs):
        calls.append((dict(kwargs['provider_snapshot']), dict(settings), dict(kwargs['existing'])))
        if len(calls) == 1:
            raise translator.TranslationCancelled(partial={segments[0].seg_id: '已完成'})
        return {**kwargs['existing'], segments[1].seg_id: '续翻完成'}
    monkeypatch.setattr(translator, 'translate_segments', translate)
    pipeline.start_translation(tid, 'p', 'auto', 'zh-CN')
    provider.update(model='new-model', base_url='http://127.0.0.1:8888/v1')
    store.save_settings({'batch_max_segments': 20, 'translate_notes': True})
    pipeline.start_translation(tid, 'p', 'auto', 'zh-CN')
    assert calls[0][0] == calls[1][0]
    assert calls[1][0]['model'] == 'first-model'
    assert calls[1][1]['batch_max_segments'] == 1
    assert calls[1][1]['translate_notes'] is False
    assert calls[1][2] == {'s000000': '已完成'}
    assert 'must-not-be-copied' not in json.dumps(store.load_job(tid))
    assert store.load_job(tid)['status'] == 'done'


def test_failed_start_persistence_releases_submission(tmp_path, monkeypatch):
    from app.services import task_manager
    monkeypatch.setattr(config, 'TASKS_DIR', tmp_path)
    tid = uuid.uuid4().hex
    store.save_job(tid, {'task_id': tid, 'status': 'pending_confirm', 'ext': '.docx'})
    def fail(*args):
        raise OSError('disk full')
    monkeypatch.setattr(store, 'save_job', fail)
    with pytest.raises(OSError, match='disk full'):
        pipeline.start_translation(tid, 'mock', 'auto', 'zh-CN')
    assert task_manager.count_in_flight() == 0


def test_deleting_default_provider_does_not_deadlock(client, tmp_path, monkeypatch):
    from app import secrets_store
    monkeypatch.setattr(config, 'PROVIDERS_FILE', tmp_path / 'providers.json')
    monkeypatch.setattr(config, 'SETTINGS_FILE', tmp_path / 'settings.json')
    monkeypatch.setattr(secrets_store, '_data_dir', lambda: tmp_path)
    monkeypatch.setattr(secrets_store, '_BACKEND', 'fernet')
    provider = client.post('/api/providers', json={'name': 'Test', 'base_url': 'http://127.0.0.1:9999/v1', 'model': 'test'}).json()['provider']
    client.put('/api/settings', json={'translation_provider_id': provider['id'], 'assistant_provider_id': provider['id']})
    assert client.delete(f'/api/providers/{provider["id"]}').json()['ok'] is True
    assert store.load_settings()['translation_provider_id'] is None
    assert store.load_settings()['assistant_provider_id'] is None
