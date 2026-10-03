"""Translated files must survive failed writes and external edits."""

from pathlib import Path

import pytest
from docx import Document

from app import config, store
from app.formats import common
from app.services import pipeline


def _task(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TASKS_DIR", tmp_path / "tasks")
    monkeypatch.setattr(config, "SETTINGS_FILE", tmp_path / "settings.json")
    task_id = "c" * 32
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("This is an English paragraph for testing")
    document.save(source)
    job = pipeline.parse_upload(task_id, source.name, source, {})
    return task_id, job


def test_external_edit_blocks_rewrite_without_losing_file(tmp_path, monkeypatch):
    task_id, job = _task(tmp_path, monkeypatch)
    seg_id = job["segments"][0]["seg_id"]
    pipeline._apply_translations(task_id, ".docx", {seg_id: "这是第一版译文"})
    exported = store.task_dir(task_id) / "translated.docx"
    stored = store.load_job(task_id)
    assert stored["revision"] == 1
    assert stored["translated_sha256"]

    external = Document(exported)
    external.add_paragraph("Edited in GenOffice")
    external.save(exported)
    before = exported.read_bytes()
    with pytest.raises(pipeline.OutputConflictError):
        pipeline.revise_segment(task_id, seg_id, "这是新的应用内译文")
    assert exported.read_bytes() == before
    assert store.load_job(task_id)["translations"][seg_id] == "这是第一版译文"
    kept = pipeline.resolve_output_conflict(task_id, "keep_external")
    assert kept["external_edit"] is True
    assert exported.read_bytes() == before
    rebuilt = pipeline.resolve_output_conflict(task_id, "rebuild")
    assert rebuilt["external_edit"] is False
    assert list((store.task_dir(task_id) / "versions").glob("external-*.docx"))
    assert "这是第一版译文" in "\n".join(p.text for p in Document(exported).paragraphs)


def test_writer_failure_keeps_last_valid_output(tmp_path, monkeypatch):
    task_id, job = _task(tmp_path, monkeypatch)
    seg_id = job["segments"][0]["seg_id"]
    pipeline._apply_translations(task_id, ".docx", {seg_id: "第一版"})
    exported = store.task_dir(task_id) / "translated.docx"
    before = exported.read_bytes()

    class BrokenWriter:
        @staticmethod
        def write_back(_source: Path, destination: Path, _translations, _options):
            destination.write_bytes(b"incomplete")
            raise RuntimeError("simulated disk failure")

    monkeypatch.setattr(common, "get_format_handler", lambda _ext: BrokenWriter)
    with pytest.raises(RuntimeError, match="simulated disk failure"):
        pipeline._apply_translations(task_id, ".docx", {seg_id: "第二版"})
    assert exported.read_bytes() == before
    assert store.load_job(task_id)["revision"] == 1
    assert not list(store.task_dir(task_id).glob(".translated-*.docx"))
