"""XLSX side-by-side preview: original on the left, translation on the right, with same-size
cells and a sheet switcher."""
import html
from pathlib import Path

import openpyxl

# Per-sheet render cap so very large workbooks cannot blow up the page.
_MAX_ROWS = 300
_MAX_COLS = 30

_STYLE = """
<style>
.xl-dual { background:#14171c; color:#e6e6e6; font-family:"Segoe UI",system-ui,sans-serif;
  padding:12px; border-radius:8px; }
.xl-tabs { display:flex; flex-wrap:wrap; gap:6px; margin-bottom:10px; }
.xl-tab { padding:4px 14px; border-radius:6px; border:1px solid #2a3040;
  background:#1d2129; color:#e6e6e6; cursor:pointer; font-size:13px; }
.xl-tab.active { background:#2a3040; font-weight:600; }
.xl-sheet { display:none; background:#1d2129; border:1px solid #2a3040; border-radius:8px;
  padding:10px; overflow:auto; }
.xl-sheet.active { display:block; }
.xl-sheet table { border-collapse:collapse; table-layout:fixed; }
.xl-sheet td, .xl-sheet th { min-width:120px; max-width:120px; border:1px solid #2a3040;
  padding:3px 6px; font-size:12px; white-space:nowrap; overflow:hidden;
  text-overflow:ellipsis; text-align:left; vertical-align:top; }
.xl-sheet th { background:#14171c; color:#9fb0c8; font-weight:600; position:sticky; top:0; }
.xl-pane { display:inline-block; vertical-align:top; margin-right:8px; margin-bottom:8px; }
.xl-pane-title { font-size:12px; color:#9fb0c8; margin-bottom:4px; }
</style>
"""

_TAB_JS = """
<script>
function xlShow(i){
  document.querySelectorAll('.xl-sheet').forEach(function(el,k){
    el.classList.toggle('active', k===i); });
  document.querySelectorAll('.xl-tab').forEach(function(el,k){
    el.classList.toggle('active', k===i); });
}
document.addEventListener('DOMContentLoaded', function(){ xlShow(0); });
</script>
"""


def _render_sheet_table(ws, title: str) -> str:
    """Render a single sheet as an HTML table (empty cells use a space placeholder so the
    grid stays aligned).

    read_only worksheets must be consumed by a single pass of iter_rows: random access via
    ws.cell() forces openpyxl to re-parse the sheet XML to locate each row, which degrades
    to O(rows^2) and can stall a request for many minutes on multi-MB xlsx files.
    """
    # Files that lack a `dimension` record (common in third-party exports) return None
    # for max_row/max_column in read_only mode — fall back to the display cap.
    n_rows = ws.max_row if isinstance(ws.max_row, int) else _MAX_ROWS
    n_cols = ws.max_column if isinstance(ws.max_column, int) else _MAX_COLS
    show_cols = min(n_cols, _MAX_COLS)

    rows_html = ["<tr><th></th>" + "".join(
        f"<th>{html.escape(openpyxl.utils.get_column_letter(c))}</th>"
        for c in range(1, show_cols + 1)) + "</tr>"]
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(n_rows, _MAX_ROWS),
                                           max_col=show_cols, values_only=True), start=1):
        # In read_only mode trailing empty cells are dropped — pad with None so columns
        # stay aligned.
        values = list(row[:show_cols]) + [None] * (show_cols - len(row))
        cells = [f"<td>{idx}</td>"]
        for v in values:
            text = "" if v is None else str(v)
            # Use a space placeholder so empty cells still hold their grid slot.
            cells.append(f'<td title="{html.escape(text, quote=True)}">'
                         f"{html.escape(text[:80]) or '&nbsp;'}</td>")
        rows_html.append("<tr>" + "".join(cells) + "</tr>")
    note = ""
    if n_rows > _MAX_ROWS or n_cols > _MAX_COLS:
        note = (f"(Showing only the first {_MAX_ROWS} rows x {_MAX_COLS} columns, "
                f"of {n_rows} rows x {n_cols} columns total)")
    return (
        f'<div class="xl-pane"><div class="xl-pane-title">{html.escape(title)} {note}</div>'
        "<table>" + "".join(rows_html) + "</table></div>"
    )


def render_xlsx_dual_html(orig_path: Path, trans_path: Path) -> str:
    """Return the full side-by-side HTML fragment (style and tab-switching JS are inlined —
    no surrounding <html> shell is needed)."""
    wo = openpyxl.load_workbook(str(orig_path), read_only=True, data_only=False)
    wt = openpyxl.load_workbook(str(trans_path), read_only=True, data_only=False)
    try:
        tabs = []
        sheets = []
        for idx, name in enumerate(wo.sheetnames):
            ws_o = wo[name]
            ws_t = wt[name] if name in wt.sheetnames else None
            tabs.append(f'<button class="xl-tab" onclick="xlShow({idx})">'
                        f"{html.escape(name)}</button>")
            panes = [_render_sheet_table(ws_o, "Original")]
            if ws_t is not None:
                panes.append(_render_sheet_table(ws_t, "Translated"))
            cls = "xl-sheet active" if idx == 0 else "xl-sheet"
            sheets.append(f'<div class="{cls}">' + "".join(panes) + "</div>")
        return (_STYLE +
                '<div class="xl-dual"><div class="xl-tabs">' + "".join(tabs) +
                "</div>" + "".join(sheets) + "</div>" + _TAB_JS)
    finally:
        wo.close()
        wt.close()
