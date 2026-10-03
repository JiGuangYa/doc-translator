"""Exercise real document commits and renderer races using controlled I/O gates."""
import io
import os
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor

import fitz
import pytest
from docx import Document

from app import config, store
from app.services import pipeline, renderer, task_manager


@pytest.fixture
def document(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'TASKS_DIR', tmp_path / 'tasks')
    monkeypatch.setattr(config, 'DATA_DIR', tmp_path)
    monkeypatch.setattr(renderer, 'auto_render_after_done', lambda *args: None)
    src = tmp_path / 'sample.docx'
    doc = Document()
    doc.add_paragraph('An original paragraph.')
    doc.save(src)
    tid = uuid.uuid4().hex
    pipeline.parse_upload(tid, src.name, src, {})
    return tid


def test_metadata_failure_preserves_revision_and_export(document, monkeypatch):
    tid = document
    pipeline._apply_translations(tid, '.docx', {'s000000': 'Saved translation.'})
    before = store.load_job(tid)
    saved = store.document_path(before, 'translated')
    previous_bytes = saved.read_bytes()
    actual_save = store.save_job
    def fail(task_id, job):
        # The new, complete file already exists when the metadata write fails.
        assert store.document_path(job, 'translated').is_file()
        raise OSError('disk full while saving job')
    monkeypatch.setattr(store, 'save_job', fail)
    with pytest.raises(OSError, match='disk full'):
        pipeline.revise_segment(tid, 's000000', 'Unsaved revision.')
    assert store.load_job(tid) == before
    assert saved.read_bytes() == previous_bytes
    assert set(p.name for p in saved.parent.glob('*.docx')) == {'original.docx', 'translated.docx'}
    monkeypatch.setattr(store, 'save_job', actual_save)
    if os.name == 'posix':
        with saved.open('rb') as reader:
            pipeline.revise_segment(tid, 's000000', 'Retried revision.')
            assert reader.read() == previous_bytes
    else:
        # Windows forbids replacing a file held by an ordinary Python reader.
        # The retained version is checked below on every platform.
        pipeline.revise_segment(tid, 's000000', 'Retried revision.')
    current = store.load_job(tid)
    assert current['content_version'] == before['content_version'] + 1
    assert Document(store.document_path(current, 'translated')).paragraphs[0].text == 'Retried revision.'
    # A reader that acquired the old job can still finish reading that version.
    assert any(path.read_bytes() == previous_bytes for path in (saved.parent / "versions").glob("*.docx"))


def test_startup_recovers_committed_file_and_keeps_version_backup(document):
    from app.services import output_transaction
    tid = document
    pipeline._apply_translations(tid, '.docx', {'s000000': 'Version one.'})
    first = store.document_path(store.load_job(tid), 'translated').read_bytes()
    pipeline.revise_segment(tid, 's000000', 'Committed version two.')
    job = store.load_job(tid)
    assert output_transaction.recover() == []
    assert store.load_job(tid) == job
    assert Document(store.document_path(job, 'translated')).paragraphs[0].text == 'Committed version two.'
    assert any(file.read_bytes() == first for file in (store.task_dir(tid) / 'versions').glob('*.docx'))


def test_start_and_revision_share_the_commit_lock(document, monkeypatch):
    tid = document
    entered, release, start_entered = threading.Event(), threading.Event(), threading.Event()
    handler = pipeline.common.get_format_handler('.docx')
    write = handler.write_back
    def paused_write(*args):
        entered.set()
        assert release.wait(5)
        return write(*args)
    monkeypatch.setattr(handler, 'write_back', paused_write)
    queued = []
    monkeypatch.setattr(task_manager, 'submit', lambda fn, tid: queued.append(fn))
    def start():
        start_entered.set()
        pipeline.start_translation(tid, 'mock', 'auto', 'en')
    with ThreadPoolExecutor(2) as pool:
        revision = pool.submit(pipeline.revise_segment, tid, 's000000', 'Manual revision.')
        assert entered.wait(5)
        starting = pool.submit(start)
        assert start_entered.wait(5)
        release.set()
        revision.result(5)
        starting.result(5)
    assert store.load_job(tid)['translations']['s000000'] == 'Manual revision.'
    # The latest revision completed the document, so no billable job is queued.
    assert queued == []


def test_cancelled_status_still_protects_final_export(document):
    tid = document
    store.update_job(tid, lambda j: j.update(status='cancelled'))
    assert task_manager.try_mark_submitted(tid)
    try:
        with pytest.raises(ValueError, match='translating'):
            pipeline.revise_segment(tid, 's000000', 'Too early')
        with pytest.raises(ValueError, match='translating'):
            store.delete_task(tid)
        assert store.load_job(tid)
    finally:
        task_manager.mark_finished(tid)


def test_resume_checkpoint_takes_precedence_over_old_placeholder(client):
    doc = Document()
    doc.add_paragraph('Original text.')
    data = io.BytesIO()
    doc.save(data)
    tid = client.post('/api/upload', files={'file': ('sample.docx', data.getvalue())}).json()['task_id']
    def update(j):
        j['segments'][0]['translation'] = '⟪ untranslated:Original text. ⟫'
        j['translations'] = {'s000000': 'New completed batch.'}
    store.update_job(tid, update)
    segment = client.get(f'/api/tasks/{tid}/segments').json()['segments'][0]
    assert segment['translation'] == 'New completed batch.'
    assert client.get(f'/api/tasks/{tid}').json()['untranslated_count'] == 0


def test_archive_restore_preserves_document_and_blocks_active_jobs(client):
    doc = Document()
    doc.add_paragraph('Original text.')
    data = io.BytesIO()
    doc.save(data)
    tid = client.post('/api/upload', files={'file': ('sample.docx', data.getvalue())}).json()['task_id']
    original = store.document_path(store.load_job(tid), 'original').read_bytes()
    response = client.patch(f'/api/tasks/{tid}', json={'archived': True})
    assert response.json()['archived'] is True
    assert client.post(f'/api/tasks/{tid}/start', json={'provider_id': 'mock'}).status_code == 400
    assert client.patch(f'/api/tasks/{tid}', json={'archived': False}).json()['archived'] is False
    assert store.document_path(store.load_job(tid), 'original').read_bytes() == original
    task_manager.try_mark_submitted(tid)
    try:
        assert client.patch(f'/api/tasks/{tid}', json={'archived': True}).status_code == 409
        assert client.delete(f'/api/tasks/{tid}').status_code == 409
        assert store.load_job(tid)['archived'] is False
    finally:
        task_manager.mark_finished(tid)


@pytest.fixture
def controlled_renderer(document, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    conversions = []
    monkeypatch.setattr(renderer, '_stopping', threading.Event())
    monkeypatch.setattr(renderer, '_threads', {})
    monkeypatch.setattr(renderer, 'genoffice_path', lambda: None)
    monkeypatch.setattr(renderer, 'soffice_path', lambda: 'fake-soffice')
    def convert(soffice, src, outdir, tag, cancel_event=None):
        conversions.append(Document(src).paragraphs[0].text)
        entered.set()
        assert release.wait(5)
        outdir.mkdir(parents=True, exist_ok=True)
        path = outdir / (src.stem + '.pdf')
        with fitz.open() as pdf:
            pdf.new_page().insert_text((40, 60), conversions[-1])
            pdf.save(path)
        return path
    monkeypatch.setattr(renderer, '_convert_to_pdf', convert)
    yield entered, release, conversions
    release.set()
    with renderer._state_lock:
        threads = list(renderer._threads.values())
    for thread in threads:
        thread.join(5)
        assert not thread.is_alive()


def drain_renderer(tid):
    # Threads can hand over to a new version before the previous thread ends.
    for _ in range(5):
        with renderer._state_lock:
            thread = renderer._threads.get(tid)
        if thread is None:
            return
        thread.join(5)
        assert not thread.is_alive()
    pytest.fail('renderer did not finish')


def test_simultaneous_preview_requests_start_one_worker(document, controlled_renderer):
    entered, release, conversions = controlled_renderer
    with ThreadPoolExecutor(8) as pool:
        responses = list(pool.map(lambda _: renderer.start_render(document, '.docx'), range(24)))
    assert entered.wait(5)
    assert all(r['status'] == 'rendering' for r in responses)
    release.set()
    drain_renderer(document)
    assert conversions == ['An original paragraph.']
    assert renderer.render_status(document)['status'] == 'ready'
    assert document not in renderer._threads


def test_revision_during_preview_discards_old_generation(document, controlled_renderer):
    entered, release, conversions = controlled_renderer
    renderer.start_render(document, '.docx')
    assert entered.wait(5)
    pipeline.revise_segment(document, 's000000', 'Newest committed revision.')
    release.set()
    drain_renderer(document)
    assert renderer.render_status(document)['status'] == 'none'
    assert renderer.page_png_path(document, 1, 'translated') is None
    renderer.start_render(document, '.docx')  # The unified reader re-requests a stale preview.
    drain_renderer(document)
    state = renderer.render_status(document)
    assert state['status'] == 'ready'
    assert state['content_version'] == store.load_job(document)['content_version'] == 1
    assert conversions[-2:] == ['An original paragraph.', 'Newest committed revision.']
    assert not (store.task_dir(document) / 'previews' / 'v0').exists()


def test_deleting_during_preview_does_not_resurrect_task(document, controlled_renderer):
    entered, release, _ = controlled_renderer
    renderer.start_render(document, '.docx')
    assert entered.wait(5)
    assert store.delete_task(document)
    release.set()
    drain_renderer(document)
    assert not store.task_dir(document).exists()
    assert document not in renderer._threads


def test_shutdown_skips_queued_preview(document, controlled_renderer):
    _, _, conversions = controlled_renderer
    with renderer._lock:
        renderer.start_render(document, '.docx')
        renderer.shutdown()
    drain_renderer(document)
    assert conversions == []
