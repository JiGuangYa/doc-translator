"""A completed model subrequest survives stopping halfway through a long paragraph."""

import hashlib

import pytest
from docx import Document

from app import store
from app.formats.common import Segment
from app.services import pipeline, task_manager
from app.services.llm import translator


def test_partial_paragraph_is_saved_and_only_unfinished_pieces_resume(tmp_path, monkeypatch):
    task_id = "c" * 32
    source = tmp_path / "long.docx"
    text = " ".join(f"Sentence {index:03d} describes a different part of the long document." for index in range(50))
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    job = pipeline.parse_upload(task_id, source.name, source, {})
    job.update(status="translating", provider_id="fake", source_lang="en", target_lang="fr")
    store.save_job(task_id, job)
    settings = {"batch_max_chars": 200, "batch_max_segments": 1,
                "concurrency_batches": 1, "use_translation_memory": False}
    seen = []

    def translate(batch, *args, **kwargs):
        seen.extend(segment.seg_id for segment in batch)
        if len(seen) == 1:
            task_manager.update(task_id, status="translating", cancel_requested=True)
        return {segment.seg_id: f"Translated {segment.seg_id}. " for segment in batch}

    monkeypatch.setattr(translator, "_translate_batch_with_retry", translate)
    pipeline._run_translation(task_id, ".docx", "fake", "en", "fr", settings)
    paused = store.load_job(task_id)
    assert paused["status"] == "cancelled"
    assert len(paused["translation_parts"]) == 1
    assert paused["translations"] == {}, "The unfinished paragraph must not be treated as translated"
    assert "translation_parts" not in store.list_job_headers()[0]

    # Recreate runtime state as at restart; the cache must come from the job file.
    task_manager._jobs.clear()
    task_manager.update(task_id, status="translating", cancel_requested=False)
    store.update_job(task_id, lambda current: current.update(status="translating"))
    pipeline._run_translation(task_id, ".docx", "fake", "en", "fr", settings)
    done = store.load_job(task_id)
    assert len(seen) > 2 and len(seen) == len(set(seen)), "Completed pieces were requested again"
    assert done["status"] == "done" and done["translation_parts"] == {}
    assert done["done_segments"] == 1 and task_manager.get_runtime(task_id)["done_segments"] == 1
    assert done["translations"]["s000000"] == "".join(f"Translated {piece}. " for piece in seen)
    assert Document(store.task_dir(task_id) / "translated.docx").paragraphs[0].text == done["translations"]["s000000"]


@pytest.mark.parametrize("invalid", ["source", "language"])
def test_changed_chunk_or_language_does_not_reuse_old_piece(monkeypatch, invalid):
    segments = [Segment(seg_id="s000000", text="Repeated source sentence. " * 30)]
    batchable, _ = translator._split_oversized(segments, 200)
    first = batchable[0]
    checkpoint = {first.seg_id: {"source_sha256": hashlib.sha256(first.text.encode()).hexdigest(),
                               "source_lang": "en", "target_lang": "fr", "translation": "Old piece"}}
    if invalid == "source":
        checkpoint[first.seg_id]["source_sha256"] = "0" * 64
    else:
        checkpoint[first.seg_id]["target_lang"] = "es"
    seen = []
    def translate(batch, *args, **kwargs):
        seen.extend(segment.seg_id for segment in batch)
        return {segment.seg_id: "New piece" for segment in batch}
    monkeypatch.setattr(translator, "_translate_batch_with_retry", translate)
    output = translator.translate_segments(segments, "fake", "en", "fr",
        {"batch_max_chars": 200, "batch_max_segments": 1}, existing_parts=checkpoint)
    assert first.seg_id in seen
    assert "Old piece" not in output["s000000"]
