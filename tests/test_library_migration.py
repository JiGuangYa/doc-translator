"""Synthetic libraries only: never read Application Support or real credentials."""

import hashlib
import os
import json

import pytest
from docx import Document
from pptx import Presentation

from app import config, store
from app.formats import common
from app.services import migration, pipeline, drafts


def write(path, value):
    store._atomic_write_json(path, value)


def legacy(tmp_path, kind, *, ext=".docx", version=1, notes=False, status="done", deleted=False):
    support = tmp_path / ("MacBook" if kind == "macbook" else "Native")
    root = support / "data" if kind == "macbook" else support
    (root / "tasks").mkdir(parents=True, exist_ok=True)
    identifier = ("2" if deleted else "1") * 32
    folder = root / ("trash" if deleted else "tasks") / identifier
    folder.mkdir(parents=True, exist_ok=True)
    original = folder / ("original" + ext)
    if ext == ".docx":
        doc = Document()
        doc.add_paragraph("Original paragraph")
        doc.add_paragraph("Another paragraph")
        doc.save(original)
        options = {"docx_version": version}
    else:
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[5])
        slide.shapes.title.text = "Slide title"
        slide.notes_slide.notes_text_frame.text = "Speaker notes"
        doc.save(original)
        options = {"translate_notes": notes}
    parsed = common.get_format_handler(ext).extract(original, options)
    segments = [
        {
            "seg_id": common.make_seg_id(i),
            "text": seg.text,
            "context": seg.context,
            "translatable": True,
            "meta": seg.meta,
        }
        for i, seg in enumerate(parsed.segments)
    ]
    translations = {segments[0]["seg_id"]: "Translated text"}
    output = folder / ("outputs/current" + ext if kind == "macbook" else "translated" + ext)
    output.parent.mkdir(exist_ok=True)
    common.get_format_handler(ext).write_back(original, output, translations, options)
    job = {
        "task_id": identifier,
        "filename": "Sample" + ext,
        "ext": ext,
        "status": status,
        "segments": segments,
        "translations": translations,
        "provider_id": "p_conflict",
        "source_lang": "en",
        "target_lang": "zh-CN",
        "archived": True,
        "content_version": 7,
        "format_options": options if kind == "native" else {},
        "source_warnings": [],
        "warnings": [],
    }
    if kind == "macbook":
        job["output_file"] = str(output.relative_to(folder))
    else:
        job["revision"] = 7
    write(folder / "job.json", job)
    write(
        root / "config/providers.json",
        {
            "providers": [
                {
                    "id": "p_conflict",
                    "name": "Fixture",
                    "model": "fixture-v1",
                    "base_url": "https://example.invalid/v1",
                }
            ]
        },
    )
    write(support / "revision-drafts.json", {identifier: {segments[0]["seg_id"]: "Unsaved draft"}})
    write(
        support / "reading-state.json",
        {
            "lastTaskID": identifier,
            "documents": {
                identifier: {
                    "mode": 1,
                    "page": 2,
                    "zoom": 1.4,
                    "paragraphID": segments[-1]["seg_id"],
                }
            },
        },
    )
    return support, root, folder, job, output


def hashes(root):
    return {
        str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in root.rglob("*")
        if p.is_file()
    }


@pytest.mark.parametrize("kind,version", [("native", 1), ("native", 2), ("macbook", 3)])
def test_copy_bytes_resume_mapping_drafts_and_idempotence(tmp_path, kind, version):
    support, root, folder, before, output = legacy(tmp_path, kind, version=version, status="paused")
    source_hashes = hashes(support)
    write(config.PROVIDERS_FILE, {"providers": [{"id": "p_conflict", "name": "Existing target"}]})
    result = migration.import_library(str(support), kind, allow_keychain=False)
    assert result["imported"] == 1 and not result["issues"]
    identifier = result["task_map"][before["task_id"]]
    job = store.load_job(identifier)
    assert job["revision"] == job["content_version"] == 7
    assert job["format_options"]["docx_version"] == version
    assert (
        job["archived"] and job["history_usage_unknown"] and job["requires_snapshot_confirmation"]
    )
    assert job["provider_id"] != "p_conflict"
    assert (
        store.document_path(store.load_job(identifier), "translated").read_bytes()
        == output.read_bytes()
    )
    assert drafts.for_task(identifier) == {"s000000": "Unsaved draft"}
    reading = json.loads((config.DATA_DIR / "reading-state.json").read_text())
    assert reading["documents"][identifier]["paragraphID"] == "s000001"
    retry = migration.import_library(str(support), kind, allow_keychain=False)
    assert retry["imported"] == 0 and retry["skipped"] == 1
    assert hashes(support) == source_hashes
    job["archived"] = False
    store.save_job(identifier, job)
    with pytest.raises(ValueError, match="确认"):
        pipeline.start_translation(identifier, job["provider_id"], "en", "zh-CN")
    pipeline.revise_segments(identifier, {"s000001": "Reviewed"}, expected_revision=7)
    assert (
        Document(store.document_path(store.load_job(identifier), "translated")).paragraphs[1].text
        == "Reviewed"
    )
    # Changed source content gets a separate copy; existing target drafts stay intact.
    old = json.loads((folder / "job.json").read_text())
    old["translations"]["s000001"] = "Later source revision"
    write(folder / "job.json", old)
    changed = migration.import_library(str(support), kind, allow_keychain=False)
    assert changed["imported"] == 1 and changed["task_map"][before["task_id"]] != identifier


@pytest.mark.parametrize("notes", [True, False])
def test_macbook_ppt_notes_inferred_and_trash_recoverable(tmp_path, notes):
    support, root, folder, before, output = legacy(
        tmp_path, "macbook", ext=".pptx", notes=notes, deleted=True
    )
    result = migration.import_library(str(support), "macbook", allow_keychain=False)
    assert not result["issues"]
    identifier = result["task_map"][before["task_id"]]
    imported = store._trash_dir() / identifier
    job = json.loads((imported / "job.json").read_text())
    assert job["format_options"]["translate_notes"] == notes
    assert (imported / "translated.pptx").read_bytes() == output.read_bytes()
    store.restore_task(identifier)
    assert store.load_job(identifier)["archived"]


def test_interrupted_import_retries_without_duplicate_or_overwriting_new_draft(
    tmp_path, monkeypatch
):
    support, root, folder, before, output = legacy(tmp_path, "macbook", version=3)
    real_write = store._atomic_write_json

    def fail(path, value):
        if path.name == "import-manifest.json" and any(
            item.get("complete") for item in value.get("tasks", {}).values()
        ):
            raise OSError("disk full")
        real_write(path, value)

    monkeypatch.setattr(store, "_atomic_write_json", fail)
    with pytest.raises(OSError):
        migration.import_library(str(support), "macbook", False)
    identifier = next(config.TASKS_DIR.iterdir()).name
    snapshot = json.loads((config.DATA_DIR / "revision-drafts.json").read_text())
    snapshot["tasks"][identifier]["segments"]["s000000"] = "Newer target edit"
    real_write(config.DATA_DIR / "revision-drafts.json", snapshot)
    monkeypatch.setattr(store, "_atomic_write_json", real_write)
    result = migration.import_library(str(support), "macbook", False)
    assert result["task_map"][before["task_id"]] == identifier
    assert len(list(config.TASKS_DIR.iterdir())) == 1
    assert drafts.for_task(identifier)["s000000"] == "Newer target edit"


def test_bad_mapping_retains_output_and_blocks_write(tmp_path):
    support, root, folder, before, output = legacy(tmp_path, "native")
    before["segments"][0]["text"] = "Not in original"
    write(folder / "job.json", before)
    result = migration.import_library(str(support), "native", False)
    identifier = result["task_map"][before["task_id"]]
    assert result["issues"]
    saved = store.document_path(store.load_job(identifier), "translated").read_bytes()
    with pytest.raises(ValueError):
        pipeline.revise_segments(identifier, {"s000000": "Wrong position"})
    assert (
        store.document_path(store.load_job(identifier), "translated").read_bytes()
        == saved
        == output.read_bytes()
    )


@pytest.mark.skipif(os.name != "posix", reason="POSIX library lease")
def test_live_source_lock_rejected(tmp_path):
    import fcntl

    support, root, *_ = legacy(tmp_path, "macbook")
    with (root / ".desktop-engine.lock").open("w") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        with pytest.raises(ValueError, match="退出"):
            migration.import_library(str(support), "macbook", False)


def test_same_machine_credentials_reencrypted_and_corrupt_credentials_reported(tmp_path):
    from cryptography.fernet import Fernet
    from app import secrets_store

    support, root, _, before, _ = legacy(tmp_path, "native")
    key = Fernet.generate_key()
    secret_dir = root / "secrets"
    secret_dir.mkdir()
    (secret_dir / ".master").write_bytes(key)
    write(
        secret_dir / "providers.json",
        {
            "p_conflict": {
                "ciphertext": Fernet(key).encrypt(b"synthetic-key").decode(),
                "keyring_current": False,
            }
        },
    )
    report = migration.import_library(str(support), "native", False)
    imported = store.load_job(report["task_map"][before["task_id"]])
    assert secrets_store.get_secret(imported["provider_id"]) == "synthetic-key"
    assert not report["credential_required"]
    # Corrupt/cross-machine ciphertext must retain configuration and require re-entry.
    provider_file = root / "config/providers.json"
    providers = json.loads(provider_file.read_text())
    providers["providers"][0]["model"] = "changed-model"
    write(provider_file, providers)
    (secret_dir / "providers.json").write_text("{damaged")
    report = migration.import_library(str(support), "native", False)
    assert len(report["credential_required"]) == 1
