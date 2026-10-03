"""User-visible recovery, revision and repeat-translation workflows."""

import copy
import os
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor

import pytest
from docx import Document

from app import config, store
from app.formats import common
from app.services import output_transaction, pipeline, task_manager


def make_task(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("A paragraph to translate and review.")
    document.save(source)
    job = pipeline.parse_upload("a" * 32, source.name, source, {})
    job.update(source_lang="en", target_lang="zh-CN", provider_id="mock")
    store.save_job(job["task_id"], job)
    return job["task_id"], job["segments"][0]["seg_id"]


def test_delete_default_provider_finishes_without_deadlock(tmp_path):
    # Isolate a potential deadlock so a regression cannot hang the whole suite.
    code = '''
from app import store
store._stored_secret_ids = lambda: set()
store.delete_secret = lambda _: None
provider = store.save_provider({"name": "Test", "base_url": "https://example.com/v1", "model": "mock"})
store.save_settings({"translation_provider_id": provider["id"]})
assert store.delete_provider(provider["id"])
assert store.get_provider(provider["id"]) is None
assert store.load_settings()["translation_provider_id"] is None
'''
    subprocess.run([sys.executable, "-c", code], check=True, timeout=3,
                   env=dict(os.environ, DOC_TRANSLATOR_DATA_DIR=str(tmp_path / "child")))


def test_stale_revision_and_concurrent_edit_do_not_overwrite(tmp_path):
    task_id, segment = make_task(tmp_path)
    pipeline._apply_translations(task_id, ".docx", {segment: "第一版"})

    def edit(text):
        try:
            pipeline.revise_segment(task_id, segment, text, expected_revision=1)
            return "saved"
        except pipeline.RevisionConflictError:
            return "conflict"

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(edit, ["第二版甲", "第二版乙"]))
    assert sorted(results) == ["conflict", "saved"]
    job = store.load_job(task_id)
    assert job["revision"] == 2
    assert Document(store.task_dir(task_id) / "translated.docx").paragraphs[0].text == job["translations"][segment]


def test_revision_conflict_api_is_distinct_and_compatible(tmp_path, client):
    task_id, segment = make_task(tmp_path)
    pipeline._apply_translations(task_id, ".docx", {segment: "第一版"})
    response = client.patch(f"/api/tasks/{task_id}/segments/{segment}",
                            json={"text": "过期草稿", "expected_revision": 0})
    assert response.status_code == 409
    assert response.json()["code"] == "revision_conflict"
    assert client.get(f"/api/tasks/{task_id}/segments").json()["revision"] == 1
    # Existing clients that do not send a revision keep their API contract.
    assert client.patch(f"/api/tasks/{task_id}/segments/{segment}", json={"text": "修订"}).status_code == 200


def test_metadata_disk_failure_rolls_back_export(tmp_path, monkeypatch):
    task_id, segment = make_task(tmp_path)
    pipeline._apply_translations(task_id, ".docx", {segment: "第一版"})
    path = store.task_dir(task_id) / "translated.docx"
    previous = path.read_bytes()
    previous_job = store.load_job(task_id)
    monkeypatch.setattr(store, "save_job", lambda *_: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError, match="disk full"):
        pipeline.revise_segment(task_id, segment, "不能保存的修订", expected_revision=1)
    assert path.read_bytes() == previous
    assert store.load_job(task_id) == previous_job
    assert not (path.parent / output_transaction.JOURNAL).exists()


def test_external_change_during_write_keeps_both_versions(tmp_path, monkeypatch):
    task_id, segment = make_task(tmp_path)
    pipeline._apply_translations(task_id, ".docx", {segment: "第一版"})
    destination = store.task_dir(task_id) / "translated.docx"
    handler = common.get_format_handler(".docx")
    original_write = handler.write_back

    def write_and_external_edit(*args):
        report = original_write(*args)
        document = Document(destination)
        document.add_paragraph("External edit while writing")
        document.save(destination)
        return report

    monkeypatch.setattr(handler, "write_back", write_and_external_edit)
    with pytest.raises(pipeline.OutputConflictError):
        pipeline.revise_segment(task_id, segment, "新的应用译文", expected_revision=1)
    assert "External edit" in Document(destination).paragraphs[-1].text
    candidate = next((destination.parent / "versions").glob("app-conflict-*.docx"))
    assert Document(candidate).paragraphs[0].text == "新的应用译文"
    assert store.load_job(task_id)["revision"] == 1


@pytest.mark.parametrize("already_replaced", [False, True])
def test_startup_recovers_interrupted_document_commit(tmp_path, already_replaced):
    task_id, _ = make_task(tmp_path)
    directory = store.task_dir(task_id)
    old = store.load_job(task_id)
    new = copy.deepcopy(old)
    new.update(revision=1, status="done")
    destination = directory / "translated.docx"
    pending = directory / ".translated-recovery.docx"
    destination.write_bytes(b"old valid output")
    pending.write_bytes(b"new validated output")
    record = {"old_job": old, "new_job": new, "destination": destination.name,
              "pending": pending.name, "old_hash": output_transaction.file_hash(destination),
              "new_hash": output_transaction.file_hash(pending)}
    store._atomic_write_json(directory / output_transaction.JOURNAL, record)
    if already_replaced:
        pending.replace(destination)
    assert output_transaction.recover() == [task_id]
    assert destination.read_bytes() == b"new validated output"
    assert store.load_job(task_id)["revision"] == 1
    assert output_transaction.recover() == []


def test_completed_model_text_survives_writer_failure(tmp_path, monkeypatch):
    task_id, segment = make_task(tmp_path)
    monkeypatch.setattr(pipeline, "_apply_translations", lambda *_: (_ for _ in ()).throw(OSError("no space")))
    with pytest.raises(OSError, match="no space"):
        pipeline._run_translation(task_id, ".docx", "mock", "en", "zh-CN", {})
    assert store.load_job(task_id)["translations"][segment]


def test_recovery_keeps_external_edit_and_proposed_app_text(tmp_path):
    task_id, segment = make_task(tmp_path)
    directory = store.task_dir(task_id)
    old = store.load_job(task_id)
    new = copy.deepcopy(old)
    new.update(revision=1, status="done", translations={segment: "待恢复的译文"})
    destination = directory / "translated.docx"
    pending = directory / ".translated-recovery.docx"
    destination.write_bytes(b"original output")
    old_hash = output_transaction.file_hash(destination)
    pending.write_bytes(b"app candidate")
    store._atomic_write_json(directory / output_transaction.JOURNAL, {
        "old_job": old, "new_job": new, "destination": destination.name,
        "pending": pending.name, "old_hash": old_hash,
        "new_hash": output_transaction.file_hash(pending)})
    destination.write_bytes(b"external edit after interruption")
    assert output_transaction.recover() == [task_id]
    assert destination.read_bytes() == b"external edit after interruption"
    recovered = store.load_job(task_id)
    assert recovered["external_edit"] is True
    assert recovered["translations"][segment] == "待恢复的译文"
    candidate = next((directory / "versions").glob("app-recovered-*.docx"))
    assert candidate.read_bytes() == b"app candidate"


def test_retarget_requires_separate_copy_and_preserves_previous(tmp_path, client):
    task_id, segment = make_task(tmp_path)
    pipeline._apply_translations(task_id, ".docx", {segment: "第一版"})
    previous = store.load_job(task_id)
    response = client.post(f"/api/tasks/{task_id}/start", json={
        "provider_id": "mock", "source_lang": "en", "target_lang": "fr"})
    assert response.status_code == 400
    clone = client.post(f"/api/tasks/{task_id}/duplicate", json={
        "source_lang": "en", "target_lang": "fr", "provider_id": "mock"})
    assert clone.status_code == 200
    new = store.load_job(clone.json()["task_id"])
    assert new["status"] == "pending_confirm" and new["target_lang"] == "fr"
    assert new["translations"] == {} and new["segments"][0]["translation"] is None
    assert store.load_job(task_id) == previous
    assert (store.task_dir(new["task_id"]) / "original.docx").read_bytes() == (store.task_dir(task_id) / "original.docx").read_bytes()


def test_queued_task_requires_cancel_before_delete_or_revise(tmp_path, client):
    task_id, segment = make_task(tmp_path)
    store.update_job(task_id, lambda job: job.update(status="queued"))
    task_manager.update(task_id, status="queued")
    assert client.delete(f"/api/tasks/{task_id}").status_code == 409
    with pytest.raises(ValueError, match="queued"):
        pipeline.revise_segment(task_id, segment, "修订")
    assert client.post(f"/api/tasks/{task_id}/cancel").json()["ok"]
    assert client.delete(f"/api/tasks/{task_id}").status_code == 200
    assert (config.TASKS_DIR.parent / "trash" / task_id / "original.docx").exists()


def test_edit_waits_for_worker_to_finish_metadata_commit(tmp_path):
    task_id, segment = make_task(tmp_path)
    pipeline._apply_translations(task_id, ".docx", {segment: "第一版"})
    task_manager.update(task_id, status="translating")
    with pytest.raises(ValueError, match="translating"):
        pipeline.revise_segment(task_id, segment, "过早修订", expected_revision=1)
    assert store.load_job(task_id)["revision"] == 1
