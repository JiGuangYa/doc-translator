"""All draft revisions are one commit, including validation and disk failures."""
import io

import fitz
import pytest
from docx import Document

from app import store
from app.services import pipeline, renderer, task_manager


@pytest.fixture
def revised_document(client, monkeypatch):
    monkeypatch.setattr(renderer, 'auto_render_after_done', lambda *args: None)
    source = Document()
    source.add_paragraph('First original paragraph.')
    source.add_paragraph('Second original paragraph.')
    data = io.BytesIO()
    source.save(data)
    tid = client.post('/api/upload', files={'file': ('bulk.docx', data.getvalue())}).json()['task_id']
    pipeline._apply_translations(tid, '.docx', {'s000000': 'Saved first.', 's000001': 'Saved second.'})
    return tid


def test_bulk_revision_commits_one_generation(client, revised_document, monkeypatch):
    tid = revised_document
    before = store.load_job(tid)['content_version']
    handler = pipeline.common.get_format_handler('.docx')
    original = handler.write_back
    writes = []
    def write(*args):
        writes.append(args[2])
        return original(*args)
    monkeypatch.setattr(handler, 'write_back', write)
    edits = {'s000000': 'Revised first.', 's000001': 'Revised second.'}
    assert client.patch(f'/api/tasks/{tid}/segments', json={'revisions': edits}).json()['ok']
    assert writes == [edits]
    assert store.load_job(tid)['content_version'] == before + 1
    output = Document(io.BytesIO(client.get(f'/api/tasks/{tid}/download').content))
    assert [p.text for p in output.paragraphs] == ['Revised first.', 'Revised second.']


@pytest.mark.parametrize('edits', [
    {}, {'s000000': 'Must not apply', 's999999': 'Unknown paragraph'}, {'../invalid': 'Invalid ID'},
])
def test_invalid_batch_never_changes_any_paragraph(client, revised_document, edits):
    tid = revised_document
    before = store.load_job(tid)
    original = client.get(f'/api/tasks/{tid}/download').content
    assert client.patch(f'/api/tasks/{tid}/segments', json={'revisions': edits}).status_code == 400
    assert store.load_job(tid) == before
    assert client.get(f'/api/tasks/{tid}/download').content == original


def test_bulk_metadata_failure_keeps_both_saved_paragraphs(client, revised_document, monkeypatch):
    tid = revised_document
    before = store.load_job(tid)
    original = client.get(f'/api/tasks/{tid}/download').content
    def fail(*args):
        raise OSError('simulated full disk')
    monkeypatch.setattr(store, 'save_job', fail)
    response = client.patch(f'/api/tasks/{tid}/segments', json={'revisions': {'s000000': 'First draft', 's000001': 'Second draft'}})
    assert response.status_code == 500
    assert store.load_job(tid) == before
    assert client.get(f'/api/tasks/{tid}/download').content == original


def test_bulk_revision_is_rejected_during_final_export(client, revised_document):
    tid = revised_document
    task_manager.try_mark_submitted(tid)
    try:
        assert client.patch(f'/api/tasks/{tid}/segments', json={'revisions': {'s000000': 'Too early'}}).status_code == 400
    finally:
        task_manager.mark_finished(tid)
    assert store.load_job(tid)['translations']['s000000'] == 'Saved first.'


def test_pdf_original_preview_before_translation(client):
    with fitz.open() as pdf:
        pdf.new_page().insert_text((50, 80), 'Preview before translation.')
        tid = client.post('/api/upload', files={'file': ('original.pdf', pdf.tobytes())}).json()['task_id']
    info = client.get(f'/api/tasks/{tid}/info').json()
    assert info['original_pages'] == 1 and info['translated_pages'] == 0
    assert info['has_translated'] is False
    image = client.get(f'/api/tasks/{tid}/preview/pdf/1?variant=original')
    assert image.content.startswith(b'\x89PNG')
