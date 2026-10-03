"""Generate large, local-only fixtures for paragraph and worksheet UI acceptance."""

import argparse
import json
import os
import sys
import uuid
from pathlib import Path

from docx import Document
from openpyxl import Workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, required=True)
    args = parser.parse_args()
    data = args.data_dir.resolve()
    production = (Path.home() / "Library/Application Support/DocTranslatorNative").resolve()
    if data == production or production in data.parents:
        raise SystemExit("Use an isolated test data directory with the test app stopped")
    root = Path(__file__).resolve().parents[2]
    os.environ["DOC_TRANSLATOR_DATA_DIR"] = str(data)
    os.environ["DOC_TRANSLATOR_DESKTOP"] = "1"
    sys.path.insert(0, str(root))
    from app import store
    from app.services import pipeline

    folder = root / "build/paging-qa"
    folder.mkdir(parents=True, exist_ok=True)
    document = Document()
    for number in range(10000):
        document.add_paragraph(f"Paragraph {number:05d}: Native paging keeps only the current page in memory. "
                               "Searching must include the whole document, including the final paragraph.")
    document.paragraphs[-1].add_run(" LAST PAGE MARKER A+B 中文")
    docx = folder / "Paging 10000 paragraphs.docx"
    document.save(docx)
    book = Workbook()
    for number, name in enumerate(["First", "Sales!2026", "末页"]):
        sheet = book.active if number == 0 else book.create_sheet()
        sheet.title = name
        for row in range(1, 151):
            sheet.cell(row, 1, f"Text row {row} in {name}")
            sheet.cell(row, 2, row)
            sheet.cell(row, 3, f"=B{row}*2")
    xlsx = folder / "Paging worksheets.xlsx"
    book.save(xlsx)
    tasks = {}
    for file in (docx, xlsx):
        task_id = uuid.uuid4().hex
        job = pipeline.parse_upload(task_id, file.name, file, {})
        job.update(provider_id="mock", source_lang="en", target_lang="zh-CN")
        store.save_job(task_id, job)
        tasks[file.name] = task_id
    (folder / "tasks.json").write_text(json.dumps(tasks, ensure_ascii=False, indent=2))
    print(json.dumps(tasks, ensure_ascii=False))


if __name__ == "__main__":
    main()
