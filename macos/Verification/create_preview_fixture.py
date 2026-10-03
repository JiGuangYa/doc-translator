"""Build and render an isolated visual fixture with more translated than source pages."""

import argparse
import json
import os
from pathlib import Path
import sys
import uuid


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", required=True)
    parser.add_argument("--report", required=True)
    args = parser.parse_args()
    data = Path(args.data_dir).resolve()
    if data.name == "DocTranslatorNative":
        raise SystemExit("Use an isolated test data directory")
    root = Path(__file__).resolve().parents[2]
    os.environ["DOC_TRANSLATOR_DATA_DIR"] = str(data)
    os.environ["DOC_TRANSLATOR_DESKTOP"] = "1"
    sys.path.insert(0, str(root))
    from docx import Document
    from app import store
    from app.services import pipeline, process_runner, renderer

    directory = root / "build/visual-fixtures"
    directory.mkdir(parents=True, exist_ok=True)
    source = directory / "Preview overflow fixture.docx"
    document = Document()
    document.add_heading("Translation pagination check", level=1)
    document.add_paragraph("This short source paragraph is deliberately expanded in the translated fixture.")
    document.save(source)
    task_id = uuid.uuid4().hex
    job = pipeline.parse_upload(task_id, source.name, source, {})
    job.update(provider_id="mock", source_lang="en", target_lang="fr")
    store.save_job(task_id, job)
    text = ("Cette page vérifie que la traduction reste entièrement accessible lorsque le texte "
            "occupe davantage de pages que le document original. ") * 110
    translations = {job["segments"][0]["seg_id"]: "Vérification de la pagination",
                    job["segments"][1]["seg_id"]: text}
    pipeline._apply_translations(task_id, ".docx", translations)
    renderer._render_task(task_id, ".docx")
    status = renderer.render_status(task_id)
    report = {"task_id": task_id, **status}
    Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2))
    assert status["status"] == "ready", status
    assert status["original_pages"] == 1 and status["translated_pages"] > 1, status
    assert process_runner.active_count() == 0
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
