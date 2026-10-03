"""Structure, package fidelity, fallback copies and persisted v1/v2 numbering."""

import copy
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from app import config, store
from app.formats import common, docx_fmt, docx_structured
from app.services import pipeline
from app.services.llm.prompts import build_user_prompt
from tests.docx_fixtures import complex_docx

OPTIONS = {"docx_version": 2}


def xml_part(path, part="word/document.xml"):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read(part))


def test_structured_identity_copies_every_byte(tmp_path):
    source = complex_docx(tmp_path / "source.docx", compatibility_copy=True)
    result = docx_fmt.extract(source, OPTIONS)
    texts = [s.text for s in result.segments]
    assert texts.count("Text inside a nested text box") == 1
    assert "Important " in texts and "read the manual" in texts
    assert "Page 1" not in texts
    output = tmp_path / "translated.docx"
    docx_fmt.write_back(source, output, {f"s{i:06d}": s.text for i, s in enumerate(result.segments)}, OPTIONS)
    assert output.read_bytes() == source.read_bytes()


@pytest.mark.parametrize("compatibility_copy", [False, True])
def test_patch_keeps_styles_links_media_and_field_parts(tmp_path, compatibility_copy):
    source = complex_docx(tmp_path / "source.docx", compatibility_copy=compatibility_copy)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("customXml/untouched.xml", b"<data>retain this extension</data>")
    result = docx_fmt.extract(source, OPTIONS)
    translations = {f"s{i:06d}": f"译文{i}" for i, _ in enumerate(result.segments)}
    output = tmp_path / "translated.docx"
    report = docx_fmt.write_back(source, output, translations, OPTIONS)
    assert report.written == len(result.segments)
    before, after = xml_part(source), xml_part(output)
    ns = {"w": docx_structured.W[1:-1]}
    assert [etree.tostring(n) for n in before.findall(".//w:rPr", ns)] == [
        etree.tostring(n) for n in after.findall(".//w:rPr", ns)]
    assert [dict(n.attrib) for n in before.findall(".//w:hyperlink", ns)] == [
        dict(n.attrib) for n in after.findall(".//w:hyperlink", ns)]
    link_text = "".join(after.find(".//w:hyperlink", ns).itertext())
    assert link_text.startswith("译文")
    boxes = after.findall(".//w:txbxContent", ns)
    assert len(boxes) == (2 if compatibility_copy else 1)
    assert len({"".join(box.itertext()) for box in boxes}) == 1
    assert "译文" in "".join(boxes[0].itertext())
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output) as translated:
        assert original.namelist() == translated.namelist()
        for name in original.namelist():
            if name not in ("word/document.xml", "word/header1.xml"):
                assert original.read(name) == translated.read(name), name
    Document(output)


def test_legacy_jobs_keep_their_numbering_and_new_copies_upgrade(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TASKS_DIR", tmp_path / "tasks")
    source = complex_docx(tmp_path / "source.docx")
    old = pipeline.parse_upload("1" * 32, source.name, source, {"docx_version": 1})
    old.pop("format_options")  # A task saved by the already installed v1 writer.
    store.save_job(old["task_id"], old)
    source_texts = [s["text"] for s in old["segments"]]
    assert "Text inside a nested text box" not in source_texts
    index = next(i for i, s in enumerate(old["segments"]) if s["text"].startswith("A final paragraph"))
    pipeline._apply_translations(old["task_id"], ".docx", {f"s{index:06d}": "旧任务仍然正确定位"})
    translated = store.task_dir(old["task_id"]) / "translated.docx"
    assert "旧任务仍然正确定位" in [p.text for p in Document(translated).paragraphs]
    clone = pipeline.duplicate_task(old["task_id"], "en", "zh-CN", "mock")
    assert clone["format_options"] == OPTIONS
    assert any(s["text"] == "Text inside a nested text box" for s in clone["segments"])
    assert next(s for s in clone["segments"] if s["text"] == "read the manual")["translatable"]
    assert store.load_job(old["task_id"])["translations"] == {f"s{index:06d}": "旧任务仍然正确定位"}


def test_default_upload_persists_structured_mapping_and_writes_link(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TASKS_DIR", tmp_path / "tasks")
    source = complex_docx(tmp_path / "source.docx")
    job = pipeline.parse_upload("2" * 32, source.name, source, {})
    assert job["format_options"] == OPTIONS
    link = next(s for s in job["segments"] if s["text"] == "read the manual")
    pipeline._apply_translations(job["task_id"], ".docx", {link["seg_id"]: "阅读手册"})
    root = xml_part(store.task_dir(job["task_id"]) / "translated.docx")
    node = root.find(".//" + docx_structured.W + "hyperlink")
    assert "".join(node.itertext()) == "阅读手册"


def test_multiple_text_nodes_tabs_and_manual_newlines_are_not_duplicated(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "target.docx"
    doc = Document()
    run = doc.add_paragraph().add_run("First ")
    extra = OxmlElement("w:t")
    extra.text = "fragment"
    run._r.append(extra)
    run.add_tab()
    run.add_text("Second fragment")
    doc.save(source)
    result = docx_fmt.extract(source, OPTIONS)
    assert [s.text for s in result.segments] == ["First fragment", "Second fragment"]
    docx_fmt.write_back(source, output, {"s000000": "第一行\n第二行", "s000001": "后续文字"}, OPTIONS)
    paragraph = Document(output).paragraphs[0]
    assert paragraph.text == "第一行\n第二行\t后续文字"


def test_non_equivalent_fallback_is_rejected(tmp_path):
    source = complex_docx(tmp_path / "source.docx", compatibility_copy=True)
    doc = Document(source)
    fallback = next(doc.element.iter(docx_structured.MC + "Fallback"))
    next(fallback.iter(docx_structured.W + "t")).text = "Different visible text"
    doc.save(source)
    with pytest.raises(common.FormatAdapterError, match="non-equivalent"):
        docx_fmt.extract(source, OPTIONS)


def test_context_sensitive_fragments_are_not_deduplicated_across_paragraphs():
    first = common.Segment("s000000", "May", meta={"context_sensitive": True, "translation_context": "In May the review starts."})
    second = common.Segment("s000001", "May", meta={"context_sensitive": True, "translation_context": "May I begin?"})
    duplicate = copy.deepcopy(first)
    duplicate.seg_id = "s000002"
    unique, aliases = common.dedupe_segments([first, second, duplicate])
    assert len(unique) == 2
    assert aliases == {"s000000": ["s000002"]}
    prompt = build_user_prompt('{"s000000":"May"}', "en", "zh-CN", contexts={first.seg_id: first.meta["translation_context"]})
    assert "<source_context>" in prompt and "In May the review starts." in prompt


def test_unknown_word_mapping_version_is_refused(tmp_path):
    source = complex_docx(tmp_path / "source.docx")
    with pytest.raises(common.FormatAdapterError, match="newer app"):
        docx_fmt.write_back(source, tmp_path / "target.docx", {"s000000": "wrong index"}, {"docx_version": 99})
    assert not (tmp_path / "target.docx").exists()


def test_legacy_multi_text_run_does_not_append_old_text(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "target.docx"
    doc = Document()
    run = doc.add_paragraph().add_run("First ")
    extra = OxmlElement("w:t")
    extra.text = "fragment"
    run._r.append(extra)
    doc.save(source)
    docx_fmt.write_back(source, output, {"s000000": "完整译文 "}, {})
    assert Document(output).paragraphs[0].text == "完整译文 "


def test_styled_fragments_do_not_use_or_pollute_context_free_memory(tmp_path, monkeypatch):
    from app.services import memory, renderer
    from app.services.llm import translator

    monkeypatch.setattr(config, "TASKS_DIR", tmp_path / "tasks")
    monkeypatch.setattr(renderer, "auto_render_after_done", lambda *args: None)
    source = tmp_path / "context.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("May").bold = True
    p.add_run(" I begin?")
    doc.add_paragraph("Standalone sentence for reuse")
    doc.save(source)
    job = pipeline.parse_upload("3" * 32, source.name, source, {})
    job.update(provider_id="fake", source_lang="en", target_lang="zh-CN")
    store.save_job(job["task_id"], job)
    looked_up, remembered, revised = [], [], []

    def lookup(source, target, texts, provider):
        looked_up.extend(texts)
        return {"May": "五月", "Standalone sentence for reuse": "可以复用的独立句子"}

    def translate(segments, *args, existing, **kwargs):
        assert "s000000" not in existing
        assert existing["s000002"] == "可以复用的独立句子"
        return {**existing, "s000000": "可以", "s000001": " 我开始吗？"}

    monkeypatch.setattr(memory, "lookup_many", lookup)
    monkeypatch.setattr(memory, "remember_many", lambda *args: remembered.extend(args[-1]))
    monkeypatch.setattr(memory, "remember", lambda *args, **kwargs: revised.append(args))
    monkeypatch.setattr(translator, "translate_segments", translate)
    pipeline._run_translation(job["task_id"], ".docx", "fake", "en", "zh-CN", {"use_translation_memory": True})
    assert looked_up == ["Standalone sentence for reuse"]
    assert remembered == [("Standalone sentence for reuse", "可以复用的独立句子")]
    pipeline.revise_segment(job["task_id"], "s000000", "能否")
    assert revised == []
