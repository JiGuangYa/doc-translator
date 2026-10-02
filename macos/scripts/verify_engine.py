"""Exercise the actual frozen engine against a local Chat Completions stub.

No external API calls or real credentials. Run with the build venv.
"""
import argparse
import io
import json
import os
import selectors
import subprocess
import threading
import time
import urllib.error
import urllib.request
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import openpyxl


class Stub(BaseHTTPRequestHandler):
    calls = 0
    models = []
    delay = 0
    expected_key = 'local-test-placeholder'
    def log_message(self, *args):
        pass
    def do_POST(self):
        if self.headers.get('Authorization') != f'Bearer {type(self).expected_key}':
            self.send_response(401); self.send_header('Content-Type', 'application/json'); self.end_headers()
            self.wfile.write(b'{"error":{"message":"Unexpected test credential"}}')
            return
        time.sleep(type(self).delay)
        body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        type(self).models.append(body['model'])
        content = body['messages'][-1]['content']
        if '{' in content:
            payload, _ = json.JSONDecoder().raw_decode(content[content.index('{'):])
            translated = {key: '这是测试译文。' for key in payload}
            content = json.dumps(translated, ensure_ascii=False)
        else:
            content = 'OK'
        type(self).calls += 1
        result = {'id': 'test-completion', 'object': 'chat.completion', 'created': 1, 'model': 'local-test',
                  'choices': [{'index': 0, 'message': {'role': 'assistant', 'content': content}, 'finish_reason': 'stop'}],
                  'usage': {'prompt_tokens': 10, 'completion_tokens': 10, 'total_tokens': 20}}
        raw = json.dumps(result).encode()
        self.send_response(200); self.send_header('Content-Type', 'application/json'); self.end_headers(); self.wfile.write(raw)


def fixtures(folder):
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document(); doc.add_heading('Reading together', 0)
    doc.add_paragraph('This document demonstrates a native bilingual reading experience.')
    doc.add_paragraph('Completed translations remain available after restarting the application.')
    doc.save(folder / 'Sample.docx')
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = 'Read and translate your documents.'
    deck.save(folder / 'Sample.pptx')
    workbook = openpyxl.Workbook(); workbook.active['A1'] = 'Document translation'; workbook.active['A2'] = 'Read both languages together'
    workbook.active['B1'] = 10; workbook.active['B2'] = '=B1+1'; workbook.save(folder / 'Sample.xlsx')
    pdf = fitz.open(); page = pdf.new_page(); page.insert_text((60, 80), 'Reading together', fontsize=24)
    page.insert_text((60, 130), 'Read the original document alongside its translation.', fontsize=12)
    pdf.save(folder / 'Sample.pdf'); pdf.close()
    with fitz.open() as long_pdf:
        for number in range(1, 21):
            page = long_pdf.new_page()
            page.insert_text((60, 80), f'Long document page {number}', fontsize=24)
            page.insert_text((60, 130), f'This paragraph belongs to page {number}. Return here after changing documents.', fontsize=12)
        long_pdf.save(folder / 'LongReading.pdf')


class Engine:
    def __init__(self, executable, data):
        self.token = uuid.uuid4().hex + uuid.uuid4().hex
        data.mkdir(parents=True, exist_ok=True)
        self.log = (data.parent / 'verification-engine.log').open('ab')
        env = dict(os.environ)
        for key in ('PYTHONPATH', 'PYTHONHOME', 'DOC_TRANSLATOR_RESOURCE_DIR', 'DOC_TRANSLATOR_DATA_DIR'):
            env.pop(key, None)
        env['PATH'] = '/usr/bin:/bin:/usr/sbin:/sbin'
        # Exercise bundled encrypted storage with fake keys, without touching
        # the user's Keychain. Keychain failure behavior is covered separately.
        env['PYTHON_KEYRING_BACKEND'] = 'keyring.backends.fail.Keyring'
        self.proc = subprocess.Popen([str(executable)], cwd=data.parent, env=env, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=self.log)
        self.proc.stdin.write(json.dumps({'token': self.token, 'data_dir': str(data)}).encode() + b'\n'); self.proc.stdin.flush()
        selector = selectors.DefaultSelector(); selector.register(self.proc.stdout, selectors.EVENT_READ)
        assert selector.select(30), 'Engine startup timed out'
        ready = json.loads(self.proc.stdout.readline()); selector.close()
        self.base = f'http://127.0.0.1:{ready["port"]}'
        for _ in range(50):
            try:
                self.request('/healthz'); break
            except urllib.error.URLError:
                time.sleep(.1)
        else:
            raise AssertionError('Engine did not become healthy')
    def request(self, path, method='GET', body=None, raw=None, content_type='application/json', authenticated=True):
        headers = {'Content-Type': content_type}
        if authenticated: headers['X-DocTranslator-Token'] = self.token
        request = urllib.request.Request(self.base + path, data=raw if raw is not None else (json.dumps(body).encode() if body is not None else None), headers=headers, method=method)
        with urllib.request.urlopen(request, timeout=40) as response:
            data = response.read()
            return json.loads(data) if 'application/json' in response.headers.get('Content-Type', '') else data
    def upload(self, file):
        boundary = uuid.uuid4().hex
        raw = f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{file.name}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode() + file.read_bytes() + f'\r\n--{boundary}--\r\n'.encode()
        return self.request('/api/upload', 'POST', raw=raw, content_type=f'multipart/form-data; boundary={boundary}')
    def close(self):
        self.proc.stdin.close()
        try:
            self.proc.wait(timeout=16)
        except subprocess.TimeoutExpired:
            self.proc.kill(); self.proc.wait(); raise AssertionError('Engine remained alive after parent pipe closed')
        self.log.close()


def wait_for(engine, path, expected, seconds=60):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        result = engine.request(path)
        if result['status'] == expected: return result
        assert result['status'] not in ('failed', 'unavailable'), result
        time.sleep(.2)
    raise AssertionError(f'Timed out: {path}: {result}')


def verify_resume(engine, executable, output, provider):
    doc = Document()
    for number in range(12):
        doc.add_paragraph(f'This is unique paragraph number {number} for interruption testing.')
    file = output / 'fixtures' / 'Resume.docx'; doc.save(file)
    tid = engine.upload(file)['task_id']
    engine.request('/api/settings', 'PUT', {'batch_max_segments': 1, 'concurrency_batches': 1})
    request = {'provider_id': provider['id'], 'source_lang': 'auto', 'target_lang': 'zh-CN'}
    Stub.delay = .2
    engine.request(f'/api/tasks/{tid}/start', 'POST', request)
    deadline = time.monotonic() + 10
    while engine.request(f'/api/tasks/{tid}')['done_segments'] < 1:
        assert time.monotonic() < deadline
        time.sleep(.05)
    engine.request(f'/api/tasks/{tid}/cancel', 'POST')
    wait_for(engine, f'/api/tasks/{tid}', 'cancelled')
    try:
        engine.request(f'/api/tasks/{tid}/start', 'POST', dict(request, target_lang='en'))
        raise AssertionError('Resume accepted another target language')
    except urllib.error.HTTPError as error:
        assert error.code == 400
    # Changing the saved provider and batching settings must not alter this task.
    engine.request(f'/api/providers/{provider["id"]}', 'PUT',
                   {'name': 'Changed configuration', 'base_url': 'http://127.0.0.1:1/v1',
                    'model': 'different-model', 'api_key': 'rotated-test-placeholder'})
    Stub.expected_key = 'rotated-test-placeholder'
    engine.request('/api/settings', 'PUT', {'batch_max_segments': 20, 'concurrency_batches': 3})
    engine.close()
    engine = Engine(executable, output / 'app' / 'data')
    engine.request(f'/api/tasks/{tid}/start', 'POST', request)
    # Close the parent pipe while translation is active, then reopen.
    time.sleep(.35)
    engine.close()
    saved = json.loads((output / 'app' / 'data' / 'tasks' / tid / 'job.json').read_text())
    completed = len(saved['translations'])
    assert 0 < completed < saved['segment_count'], saved
    engine = Engine(executable, output / 'app' / 'data')
    before_calls = Stub.calls
    Stub.delay = 0
    engine.request(f'/api/tasks/{tid}/start', 'POST', request)
    result = wait_for(engine, f'/api/tasks/{tid}', 'done')
    assert result['untranslated_count'] == 0
    assert Stub.calls - before_calls == saved['segment_count'] - completed
    assert set(Stub.models) == {'local-test'}, Stub.models
    assert result['provider_snapshot']['model'] == 'local-test'
    print('Cancel, active quit, restart, and resume with the original model, endpoint and settings passed', flush=True)
    return engine


def main():
    parser = argparse.ArgumentParser(); parser.add_argument('engine', type=Path); parser.add_argument('output', type=Path)
    args = parser.parse_args(); output = args.output.resolve(); output.mkdir(parents=True, exist_ok=True)
    fixtures(output / 'fixtures')
    server = ThreadingHTTPServer(('127.0.0.1', 0), Stub)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    engine = Engine(args.engine.resolve(), output / 'app' / 'data')
    results = []
    try:
        duplicate = subprocess.run([str(args.engine.resolve())], input=json.dumps({
            'token': uuid.uuid4().hex * 2, 'data_dir': str(output / 'app' / 'data')
        }).encode() + b'\n', stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=10)
        assert duplicate.returncode == 0
        assert '另一份' in json.loads(duplicate.stdout)['error']
        assert engine.request('/healthz')['ok']
        try:
            engine.request('/api/tasks', authenticated=False)
            raise AssertionError('Unauthenticated API was accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 401
        provider = engine.request('/api/providers', 'POST', {'name': '本地验收模拟 API', 'base_url': f'http://127.0.0.1:{server.server_port}/v1', 'model': 'local-test', 'api_key': Stub.expected_key})['provider']
        assert provider['has_api_key'] is True
        assert engine.request(f'/api/providers/{provider["id"]}/test', 'POST')['ok']
        for ext in ('docx', 'pptx', 'xlsx', 'pdf'):
            task = engine.upload(output / 'fixtures' / f'Sample.{ext}'); tid = task['task_id']
            if ext == 'pdf':
                original = engine.request(f'/api/tasks/{tid}/info')
                assert original['original_pages'] > 0 and original['translated_pages'] == 0
                original_image = engine.request(f'/api/tasks/{tid}/preview/pdf/1?variant=original')
            else:
                engine.request(f'/api/tasks/{tid}/preview/render', 'POST')
                original = wait_for(engine, f'/api/tasks/{tid}/preview/render', 'ready', 90)
                assert original['original_pages'] > 0 and original['translated_pages'] == 0
                original_image = engine.request(f'/api/tasks/{tid}/preview/render/page/1?variant=original')
            assert original_image.startswith(b'\x89PNG')
            engine.request(f'/api/tasks/{tid}/start', 'POST', {'provider_id': provider['id'], 'source_lang': 'auto', 'target_lang': 'zh-CN'})
            done = wait_for(engine, f'/api/tasks/{tid}', 'done')
            assert done['untranslated_count'] == 0, done
            segments = engine.request(f'/api/tasks/{tid}/segments')['segments']
            assert all(segment['translation'] == '这是测试译文。' for segment in segments if segment['translatable'])
            seg = next(segment for segment in segments if segment['translatable'])
            engine.request(f'/api/tasks/{tid}/segments/{seg["seg_id"]}', 'PATCH', {'text': 'Revised translation.'})
            revisions = {segment['seg_id']: ('Revised translation.' if segment['seg_id'] == seg['seg_id'] else 'Reviewed paragraph.')
                         for segment in segments if segment['translatable']}
            engine.request(f'/api/tasks/{tid}/segments', 'PATCH', {'revisions': revisions})
            assert engine.request(f'/api/tasks/{tid}')['content_version'] == 3
            exported = engine.request(f'/api/tasks/{tid}/download')
            export_path = output / f'Revised.{ext}'; export_path.write_bytes(exported)
            reopened = engine.upload(export_path)
            extracted = engine.request(f'/api/tasks/{reopened["task_id"]}/segments')['segments']
            assert any('Revised translation.' in segment['text'] for segment in extracted), extracted
            if ext == 'pdf':
                info = engine.request(f'/api/tasks/{tid}/info')
                assert info['original_pages'] == info['translated_pages'] == 1
                image = engine.request(f'/api/tasks/{tid}/preview/pdf/1?variant=translated')
            else:
                engine.request(f'/api/tasks/{tid}/preview/render', 'POST')
                rendered = wait_for(engine, f'/api/tasks/{tid}/preview/render', 'ready', 90)
                assert rendered['content_version'] == 3, rendered
                assert rendered['original_pages'] > 0 and rendered['translated_pages'] > 0
                image = engine.request(f'/api/tasks/{tid}/preview/render/page/1?variant=translated')
            assert image.startswith(b'\x89PNG')
            (output / f'Preview-{ext}.png').write_bytes(image)
            assert engine.request(f'/api/tasks/{tid}', 'PATCH', {'archived': True})['archived']
            assert engine.request(f'/api/tasks/{tid}/download') == exported
            assert not engine.request(f'/api/tasks/{tid}', 'PATCH', {'archived': False})['archived']
            assert engine.request(f'/api/tasks/{tid}/download') == exported
            if ext == 'docx':
                engine.request(f'/api/tasks/{tid}', 'PATCH', {'archived': True})
            results.append({'format': ext, 'translated': True, 'revised': True, 'bulk_revision': True, 'original_preview': True, 'reopened': True, 'preview': True, 'archive_restore': True, 'task_id': tid})
            print(f'{ext}: translate → revise → export → reopen → preview passed', flush=True)
        engine = verify_resume(engine, args.engine.resolve(), output, provider)
        engine.request('/api/settings', 'PUT', {'translation_provider_id': provider['id']})
        assert engine.request(f'/api/providers/{provider["id"]}', 'DELETE')['ok']
        assert engine.request('/api/providers')['settings']['translation_provider_id'] is None
        secret_records = json.loads((output / 'app/data/secrets/providers.json').read_text())
        assert provider['id'] not in secret_records
        ids = {task['task_id'] for task in engine.request('/api/tasks')}
        archived_ids = {task['task_id'] for task in engine.request('/api/tasks') if task['archived']}
    finally:
        engine.close()
    # Restart with a new private token and verify durable history.
    restarted = Engine(args.engine.resolve(), output / 'app' / 'data')
    try:
        assert ids == {task['task_id'] for task in restarted.request('/api/tasks')}
        assert archived_ids == {task['task_id'] for task in restarted.request('/api/tasks') if task['archived']}
        for task in restarted.request('/api/tasks'):
            if task['has_translated']:
                folder = output / 'app' / 'data' / 'tasks' / task['task_id']
                record = json.loads((folder / 'job.json').read_text())
                assert (folder / record['output_file']).is_file()
                assert len(list((folder / 'outputs').iterdir())) == 1
                assert restarted.request(f'/api/tasks/{task["task_id"]}/download')
        assert restarted.token != engine.token
    finally:
        restarted.close(); server.shutdown()
    (output / 'engine-verification.json').write_text(json.dumps({'results': results, 'api_calls': Stub.calls, 'restart': True, 'parent_exit_cleanup': True, 'cancel_and_resume': True, 'active_quit_resume': True, 'frozen_configuration': True, 'encrypted_key_rotation': True, 'delete_default_provider': True, 'single_writer': True, 'archive_restore': True, 'committed_output_recovery': True}, ensure_ascii=False, indent=2))
    print('Restart and parent-exit cleanup passed', flush=True)


if __name__ == '__main__':
    main()
